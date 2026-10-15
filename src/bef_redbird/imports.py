"""Bounded, values-only Excel adapters. Originals are never modified or extracted."""

from __future__ import annotations

import hashlib
import io
import re
import struct
from collections.abc import Iterator
from dataclasses import replace
from datetime import date, datetime
from typing import Any
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

import olefile
import xlrd
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from xlrd.biffh import XLRDError

from bef_redbird.domain import (
    CitoItem,
    CitoSnapshot,
    DataError,
    HistoryDataset,
    Issue,
    SourceInfo,
)
from bef_redbird.validation import group_id, normal_header, price_ore, product_id

MAX_UPLOAD_BYTES = 128 * 1024 * 1024
MAX_EXPANDED_BYTES = 1024 * 1024 * 1024
MAX_CELLS = 25_000_000
MAX_ROWS = 250_000
MAX_COLUMNS = 600
REQUIRED_CITO = ("varenr", "opbevkode", "kostpris", "subkode", "nysubkode", "subgrp")


def source_info(data: bytes, name: str) -> SourceInfo:
    if not data or len(data) > MAX_UPLOAD_BYTES:
        raise DataError("Filen er tom eller større end 128 MB.")
    basename = name.replace("\\", "/").rsplit("/", 1)[-1][:200]
    return SourceInfo(basename, hashlib.sha256(data).hexdigest(), len(data))


def _xlsx(data: bytes, name: str) -> Any:
    source_info(data, name)
    if not name.lower().endswith(".xlsx"):
        raise DataError("Brug en .xlsx-fil uden makroer eller kryptering.")
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            names = {i.filename for i in entries}
            if len(entries) > 5000 or sum(i.file_size for i in entries) > MAX_EXPANDED_BYTES:
                raise DataError("Excel-filen overskrider grænsen for udpakket størrelse.")
            if {"[Content_Types].xml", "xl/workbook.xml"} - names:
                raise DataError("Filen er ikke en understøttet Excel-projektmappe.")
            if any(i.flag_bits & 1 or i.file_size > MAX_EXPANDED_BYTES for i in entries):
                raise DataError("Krypteret eller for stor Excel-fil.")
            if any("vbaproject" in n.lower() or "macrosheets/" in n.lower() for n in names):
                raise DataError("Makroindhold understøttes ikke.")
            content_types = archive.read("[Content_Types].xml")
            if b"macroEnabled" in content_types:
                raise DataError("Makroindhold understøttes ikke.")
        return load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except (
        BadZipFile,
        InvalidFileException,
        KeyError,
        DefusedXmlException,
        OSError,
        ParseError,
        ValueError,
    ) as exc:
        if isinstance(exc, DataError):
            raise
        raise DataError("Excel-filen kan ikke læses. Brug en ukrypteret værdieksport.") from exc


def _rows(sheet: Any) -> Iterator[tuple[int, list[Any]]]:
    # Worksheet dimensions and defined names are not trusted as source boundaries.
    sheet.reset_dimensions()
    cells = 0
    try:
        for index, row in enumerate(sheet.iter_rows(), 1):
            cells += len(row)
            if index > MAX_ROWS or len(row) > MAX_COLUMNS or cells > MAX_CELLS:
                raise DataError("Arket overskrider grænsen for rækker, kolonner eller celler.")
            yield index, list(row)
    except (ParseError, DefusedXmlException, ValueError, KeyError) as exc:
        if isinstance(exc, DataError):
            raise
        raise DataError(
            "Arket indeholder ugyldig Excel-struktur; brug en ny værdieksport."
        ) from exc


def _header_map(values: list[Any], required: tuple[str, ...]) -> dict[str, int]:
    result: dict[str, int] = {}
    for col, value in enumerate(values):
        header = normal_header(value)
        if header in required:
            if header in result:
                raise DataError("To kolonner har samme påkrævede overskrift.")
            result[header] = col
    if set(required) - result.keys():
        raise DataError("Påkrævede kolonner mangler. Kontroller eksportens overskrifter.")
    return result


def _biff_formula_cells(data: bytes) -> dict[int, set[tuple[int, int]]]:
    """Inspect BIFF8 FORMULA records before xlrd can expose cached results."""
    try:
        with olefile.OleFileIO(io.BytesIO(data)) as ole:
            paths = ole.listdir()
            if any(
                any(p.lower() in ("vba", "_vba_project_cur", "macros") for p in path)
                for path in paths
            ):
                raise DataError("Makroindhold understøttes ikke.")
            stream_name = "Workbook" if ole.exists("Workbook") else "Book"
            if not ole.exists(stream_name):
                raise DataError("Ugyldig ældre Excel-fil.")
            stream = ole.openstream(stream_name).read()
        offset, current_sheet = 0, -1
        offsets: dict[int, int] = {}
        formulas: dict[int, set[tuple[int, int]]] = {}
        while offset + 4 <= len(stream):
            record, size = struct.unpack_from("<HH", stream, offset)
            payload = stream[offset + 4 : offset + 4 + size]
            if len(payload) != size:
                raise DataError("Afkortet ældre Excel-fil.")
            if record == 0x002F:
                raise DataError("Krypterede Excel-filer understøttes ikke.")
            if record == 0x0085 and size >= 6:  # BOUNDSHEET
                if payload[5] not in (0, 2):
                    raise DataError("Makroark understøttes ikke.")
                # xlrd's worksheet indexes exclude chart sheets. Keep the same mapping
                # so a chart before the data sheet cannot hide its FORMULA records.
                if payload[5] == 0:
                    offsets[struct.unpack_from("<I", payload)[0]] = len(offsets)
            if record == 0x0809 and size >= 4:
                if struct.unpack_from("<H", payload, 2)[0] == 0x0040:
                    raise DataError("Makroark understøttes ikke.")
                current_sheet = offsets.get(offset, -1)
            if record == 0x0006 and size >= 4 and current_sheet >= 0:
                row, col = struct.unpack_from("<HH", payload)
                formulas.setdefault(current_sheet, set()).add((row, col))
            offset += size + 4
        return formulas
    except (OSError, struct.error) as exc:
        raise DataError("Den ældre Excel-fil kan ikke læses sikkert.") from exc


def _cito_rows(data: bytes, name: str) -> tuple[str, list[tuple[int, list[Any]]]]:
    # Cito is small. Only the six required cells are retained, never optional private fields.
    if name.lower().endswith(".xls"):
        formulas = _biff_formula_cells(data)
        try:
            book = xlrd.open_workbook(file_contents=data, on_demand=True, logfile=io.StringIO())
        except XLRDError as exc:
            raise DataError("Den ældre Excel-fil kan ikke læses.") from exc
        try:
            matches: list[tuple[Any, dict[str, int], int]] = []
            for i, sheet in enumerate(book.sheets()):
                if sheet.nrows > MAX_ROWS or sheet.ncols > MAX_COLUMNS:
                    raise DataError("Cito-arket overskrider størrelsesgrænsen.")
                if not sheet.nrows:
                    continue
                headers = [normal_header(v) for v in sheet.row_values(0)]
                if set(REQUIRED_CITO) <= set(headers):
                    matches.append(
                        (sheet, _header_map(list(sheet.row_values(0)), REQUIRED_CITO), i)
                    )
            if len(matches) != 1:
                raise DataError(
                    "Ældre Cito-fil skal have præcis ét ark med de seks påkrævede felter."
                )
            sheet, mapping, index = matches[0]
            required_cols = set(mapping.values())
            if any(col in required_cols for _, col in formulas.get(index, set())):
                raise DataError("Formler i påkrævede Cito-felter: brug en værdieksport.")
            result = []
            for row_index in range(1, sheet.nrows):
                values = []
                for field in REQUIRED_CITO:
                    cell = sheet.cell(row_index, mapping[field])
                    values.append(
                        "#INVALID_CELL" if cell.ctype == xlrd.XL_CELL_ERROR else cell.value
                    )
                result.append((row_index + 1, values))
            return sheet.name, result
        finally:
            book.release_resources()
    book = _xlsx(data, name)
    try:
        if "Table" not in book.sheetnames:
            raise DataError("Cito-filen mangler arket Table.")
        rows = _rows(book["Table"])
        _, header = next(rows, (0, []))
        mapping = _header_map([c.value for c in header], REQUIRED_CITO)
        result = []
        for rowno, row in rows:
            values = []
            for field in REQUIRED_CITO:
                col = mapping[field]
                cell = row[col] if col < len(row) else None
                if cell is not None and cell.data_type == "f":
                    raise DataError("Formler i påkrævede Cito-felter: brug en værdieksport.")
                values.append(
                    ("#INVALID_CELL" if cell.data_type == "e" else cell.value)
                    if cell is not None
                    else None
                )
            result.append((rowno, values))
        return "Table", result
    finally:
        book.close()


def parse_cito(data: bytes, name: str = "cito.xlsx") -> CitoSnapshot:
    source = source_info(data, name)
    sheet_name, rows = _cito_rows(data, name)
    grouped: dict[str, dict[tuple[Any, ...], CitoItem]] = {}
    issues: list[Issue] = []
    invalid_ids: set[str] = set()
    incomplete: set[str] = set()
    count, duplicates = 0, 0
    for rowno, values in rows:
        if not any(v is not None and str(v).strip() for v in values):
            continue
        count += 1
        pid: str | None = None
        gid: str | None = None
        group_valid = True
        try:
            gid = group_id(values[5])
        except DataError:
            group_valid = False
        try:
            pid = product_id(values[0])
            if not group_valid or "#INVALID_CELL" in values:
                raise DataError("Excel-fejl i påkrævede Cito-felter.")
            storage = str(values[1] or "").strip().upper()
            cost = price_ore(values[2])
            current, upcoming = (str(values[i] or "").strip().upper() for i in (3, 4))
            # Blank classifications are valid source missingness, never eligibility.
            if (
                cost is None
                or len(storage) > 16
                or current not in ("", "A", "B", "C")
                or upcoming not in ("", "A", "B", "C")
            ):
                raise DataError("Påkrævede værdier mangler eller er ugyldige.")
            key = (pid, storage, cost, current, upcoming, gid)
            bucket = grouped.setdefault(pid, {})
            if key in bucket:
                duplicates += 1
                bucket[key] = replace(bucket[key], source_rows=(*bucket[key].source_rows, rowno))
            else:
                bucket[key] = CitoItem(pid, storage, cost, current, upcoming, gid, (rowno,))
        except DataError:
            issues.append(
                Issue("invalid_cito", "Ugyldige påkrævede Cito-værdier.", "unranked", pid, rowno)
            )
            if pid:
                invalid_ids.add(pid)
            if gid:
                incomplete.add(gid)
    items: dict[str, CitoItem] = {}
    conflicts = set(invalid_ids)
    for pid, bucket in grouped.items():
        if len(bucket) > 1 or pid in invalid_ids:
            conflicts.add(pid)
            incomplete.update(v.group_id for v in bucket.values() if v.group_id is not None)
            for item in bucket.values():
                issues.append(
                    Issue(
                        "cito_conflict",
                        "Modstridende Cito-rækker; produktet udeladt fra beregning.",
                        "unranked",
                        pid,
                        item.source_rows[0],
                    )
                )
        else:
            items[pid] = next(iter(bucket.values()))
    if len(issues) > 10000:
        raise DataError("For mange ugyldige rækker; kontroller Cito-eksporten.")
    return CitoSnapshot(
        source,
        items,
        issues,
        count,
        sum(len(v) for v in grouped.values()),
        duplicates,
        conflicts,
        incomplete,
        sheet_name,
    )


def _period(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if re.fullmatch(r"[0-9]{8}", text):
        try:
            return datetime.strptime(text, "%Y%m%d").date()
        except ValueError as exc:
            raise DataError("Ugyldig datooverskrift i LMS-arket.") from exc
    if re.match(r"^[0-9]{4}[-/.][0-9]", text) or re.fullmatch(r"[0-9]{6,}", text):
        raise DataError("Datooverskrifter skal være Excel-datoer eller YYYYMMDD.")
    return None


def parse_lms_history(data: bytes, name: str = "history.xlsx") -> HistoryDataset:
    source = source_info(data, name)
    book = _xlsx(data, name)
    issues: list[Issue] = []
    prices: dict[str, tuple[int | None, ...]] = {}
    labels: dict[str, str] = {}
    source_rows: dict[str, tuple[int, ...]] = {}
    conflicts: set[str] = set()
    conflicted_periods: dict[str, set[date]] = {}
    try:
        if "Medicinpriser data" not in book.sheetnames:
            raise DataError("LMS-filen mangler arket Medicinpriser data.")
        rows = _rows(book["Medicinpriser data"])
        _, header = next(rows, (0, []))
        values = [c.value for c in header]
        mapping = _header_map(values, ("varenummer", "indikator"))
        dates: dict[date, int] = {}
        known = {
            "atc",
            "lægemiddel",
            "indholdsstof",
            "varenummer",
            "pakning",
            "styrke",
            "form",
            "firma",
            "indikator",
        }
        for col, cell in enumerate(header):
            d = _period(cell.value)
            if d:
                if d in dates:
                    raise DataError("Gentagne datooverskrifter er tvetydige; ret eksporten.")
                dates[d] = col
        if not dates:
            raise DataError("Ingen prisperioder fundet i LMS-arket.")
        # Unknown reference columns outside the date block are optional metadata.
        # Inside it, a blank/unknown heading could hide an observation; reject it.
        # Known reference fields remain reorderable, even between date columns.
        first_date_col, last_date_col = min(dates.values()), max(dates.values())
        date_cols = set(dates.values())
        for col in range(first_date_col, last_date_col + 1):
            if col not in date_cols and normal_header(values[col]) not in known:
                raise DataError("Ugyldig overskrift i LMS-arkets datokolonner.")
        periods = tuple(sorted(dates))
        for first, second in zip(periods, periods[1:], strict=False):
            if (second - first).days != 14:
                issues.append(
                    Issue(
                        "irregular_calendar",
                        "Priskalenderen afviger fra 14 dage; tvetydige naboberegninger udelades.",
                    )
                )
        label_col = next(
            (i for i, v in enumerate(values) if normal_header(v) == "lægemiddel"), None
        )
        for rowno, row in rows:

            def cell_at(col: int, row: list[Any] = row) -> Any:
                return row[col] if col < len(row) else None

            indicator = cell_at(mapping["indikator"])
            if indicator is not None and indicator.data_type == "f":
                raise DataError("Formel i Indikator: brug en værdieksport.")
            if indicator is None or indicator.value != "AIP":
                continue
            id_cell = cell_at(mapping["varenummer"])
            if id_cell is not None and id_cell.data_type == "f":
                raise DataError("Formel i Varenummer: brug en værdieksport.")
            try:
                pid = product_id(id_cell.value if id_cell is not None else None)
            except DataError:
                issues.append(
                    Issue("invalid_history_id", "Ugyldigt LMS-varenummer.", "unranked", row=rowno)
                )
                continue
            series: list[int | None] = []
            for d in periods:
                cell = cell_at(dates[d])
                if cell is not None and cell.data_type == "f":
                    raise DataError("Formel i AIP-pris: brug en værdieksport.")
                try:
                    value = price_ore(cell.value if cell is not None else None)
                    if value == 0:
                        issues.append(
                            Issue(
                                "zero_aip",
                                "AIP er nul og kan ikke indgå i beregning.",
                                "warning",
                                pid,
                                rowno,
                                d.isoformat(),
                            )
                        )
                    series.append(value)
                except DataError:
                    series.append(None)
                    issues.append(
                        Issue(
                            "invalid_aip",
                            "Ugyldig AIP; observationen indgår ikke i beregning.",
                            "warning",
                            pid,
                            rowno,
                            d.isoformat(),
                        )
                    )
            converted = tuple(series)
            if pid in prices:
                ambiguous = conflicted_periods.get(pid, set())
                for period, old, new in zip(periods, prices[pid], converted, strict=True):
                    if old != new and period not in ambiguous:
                        ambiguous.add(period)
                        issues.append(
                            Issue(
                                "history_conflict",
                                "Modstridende AIP i perioden; observationen er udeladt.",
                                "warning",
                                pid,
                                rowno,
                                period.isoformat(),
                            )
                        )
                if ambiguous:
                    conflicts.add(pid)
                    conflicted_periods[pid] = ambiguous
                    prices[pid] = tuple(
                        None if period in ambiguous else value
                        for period, value in zip(periods, prices[pid], strict=True)
                    )
            else:
                prices[pid] = converted
            source_rows[pid] = (*source_rows.get(pid, ()), rowno)
            if label_col is not None:
                label = cell_at(label_col)
                if label is not None and label.data_type not in ("f", "e"):
                    labels[pid] = str(label.value or "")[:200]
            if len(issues) > 10000:
                raise DataError("For mange ugyldige AIP-værdier; kontroller LMS-eksporten.")
        return HistoryDataset(
            source, periods, prices, labels, source_rows, issues, conflicts, conflicted_periods
        )
    finally:
        book.close()
