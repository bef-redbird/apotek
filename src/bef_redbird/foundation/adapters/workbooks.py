"""Versioned workbook assertions, independent of application row reconciliation.

Only bounded container/worksheet readers are shared with the application. Every
source row remains an assertion; dates, exact money and presence are interpreted
here without selecting a preferred duplicate or backdating current attributes.
"""

from __future__ import annotations

import io
from collections.abc import Generator
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

import xlrd
from openpyxl.utils.cell import get_column_letter
from xlrd.biffh import XLRDError

from bef_redbird.domain import DataError
from bef_redbird.imports import (
    MAX_CELLS,
    MAX_COLUMNS,
    MAX_ROWS,
    REQUIRED_CITO,
    _biff_formula_cells,
    _header_map,
    _period,
    _rows,
    _xlsx,
    source_info,
)
from bef_redbird.validation import group_id, normal_header

from .base import AdapterError, Assertion, Parsed, Problem, field_state, identifier, money

CITO_VERSION = "cito-workbook-v1"
HISTORY_VERSION = "lms-history-workbook-v1"
_PACKAGE_FIELDS = ("name", "form", "strength", "pack_description", "company", "storage_code", "atc")
_HISTORY_FIELDS = {
    "lægemiddel": "name",
    "form": "form",
    "styrke": "strength",
    "pakning": "pack_description",
    "firma": "company",
    "atc": "atc",
}


@dataclass(frozen=True)
class _Cell:
    value: Any = None
    kind: str = "n"


def _raw(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    # Represent numeric values as text too: a BIFF 123.0 stays distinguishable
    # from a string identifier and from its canonical 000123 mapping.
    if isinstance(value, (float, int)) and not isinstance(value, bool):
        return str(value)
    return value


def _at(row: list[_Cell], col: int) -> _Cell:
    return row[col] if col < len(row) else _Cell()


def _locator(sheet: str, row: int, col: int | None = None) -> str:
    suffix = f"row:{row}" if col is None else f"{get_column_letter(col + 1)}{row}"
    return f"sheet:{sheet}!{suffix}"


def _context_period(context: dict[str, Any], key: str, *, required: bool = False) -> date | None:
    value = context.get(key)
    if value is None and not required:
        return None
    try:
        if not isinstance(value, str):
            raise ValueError
        parsed = date.fromisoformat(value)
        if parsed.isoformat() != value:
            raise ValueError
        return parsed
    except ValueError as exc:
        raise AdapterError(f"invalid_{key}_context") from exc


def _price_context(context: dict[str, Any], *, cito: bool) -> dict[str, Any]:
    required: tuple[str, ...] = ("tax_basis", "fee_basis")
    if cito:
        required += ("price_type", "currency", "price_basis")
    if any(not isinstance(context.get(k), str) or not context[k].strip() for k in required):
        raise AdapterError("missing_price_context")
    if cito and (
        context["price_type"] not in ("aip", "registered", "consumer")
        or context["currency"] != "DKK"
        or context["price_basis"] not in ("per_package", "per_unit")
    ):
        raise AdapterError("invalid_price_context")
    if not cito and any(
        key in context and context[key] != expected
        for key, expected in (
            ("price_type", "aip"),
            ("currency", "DKK"),
            ("price_basis", "per_package"),
        )
    ):
        raise AdapterError("conflicting_history_price_context")
    return {
        "price_type": context["price_type"] if cito else "aip",
        "currency": "DKK",
        "price_basis": context["price_basis"] if cito else "per_package",
        "tax_basis": context["tax_basis"],
        "fee_basis": context["fee_basis"],
        "price_role": "effective_price",
    }


def _xlsx_rows(data: bytes, sheet: str) -> Generator[tuple[int, list[_Cell]], None, None]:
    book = _xlsx(data, "source.xlsx")
    try:
        if sheet not in book.sheetnames:
            raise AdapterError("missing_required_sheet")
        for rowno, cells in _rows(book[sheet]):
            yield rowno, [_Cell(c.value, c.data_type) for c in cells]
    finally:
        book.close()


def _cito_rows(data: bytes) -> tuple[str, dict[str, int], list[tuple[int, list[_Cell]]]]:
    if data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        source_info(data, "source.xls")
        formulas = _biff_formula_cells(data)
        book = xlrd.open_workbook(file_contents=data, on_demand=True, logfile=io.StringIO())
        try:
            matches = []
            for index, sheet in enumerate(book.sheets()):
                if (
                    sheet.nrows > MAX_ROWS
                    or sheet.ncols > MAX_COLUMNS
                    or sheet.nrows * sheet.ncols > MAX_CELLS
                ):
                    raise AdapterError("workbook_resource_limit")
                if sheet.nrows and set(REQUIRED_CITO) <= {
                    normal_header(v) for v in sheet.row_values(0)
                }:
                    matches.append((sheet, index))
            if len(matches) != 1:
                raise AdapterError("ambiguous_cito_sheet")
            sheet, index = matches[0]
            _validate_headers(
                [
                    _Cell(
                        cell.value,
                        "f"
                        if (0, col) in formulas.get(index, set())
                        else "e"
                        if cell.ctype == xlrd.XL_CELL_ERROR
                        else "n",
                    )
                    for col, cell in enumerate(sheet.row(0))
                ]
            )
            mapping = _header_map(list(sheet.row_values(0)), REQUIRED_CITO)
            rows = []
            for rowno in range(1, sheet.nrows):
                cells = []
                for col in range(sheet.ncols):
                    cell = sheet.cell(rowno, col)
                    kind = "f" if (rowno, col) in formulas.get(index, set()) else "n"
                    value = cell.value
                    if kind == "f":
                        pass
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        kind = "e"
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        kind, value = "b", bool(value)
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        # A formatted Excel date is not an ID or amount. Retain its
                        # original serial and explicit type rather than reinterpret it
                        # as numeric business evidence (or invent a calendar conversion).
                        kind = "d"
                    cells.append(_Cell(value, kind))
                rows.append((rowno + 1, cells))
            return sheet.name, mapping, rows
        finally:
            book.release_resources()
    rows_iter = _xlsx_rows(data, "Table")
    try:
        _, header = next(rows_iter, (0, []))
        _validate_headers(header)
        mapping = _header_map([c.value for c in header], REQUIRED_CITO)
        return "Table", mapping, list(rows_iter)
    finally:
        rows_iter.close()


def _validate_headers(header: list[_Cell]) -> None:
    if any(cell.kind in ("f", "e") for cell in header):
        raise AdapterError("invalid_header_cell")


def _field(
    cell: _Cell,
    parsed: Parsed,
    locator: str,
    pid: str,
    field: str,
    period: date | None,
    *,
    conversion: str = "text",
) -> tuple[Any, str]:
    state = field_state(cell.value)
    code = "formula_cell" if cell.kind == "f" else "error_cell"
    try:
        if cell.kind in ("f", "e"):
            raise ValueError(code)
        if cell.kind == "d":
            raise ValueError("invalid_date_cell")
        if state == "blank":
            return None, state
        if conversion == "money":
            return money(cell.value), state
        if conversion == "group":
            return group_id(cell.value), state
        value = str(cell.value).strip()
        if conversion == "abc":
            value = value.upper()
            if value not in ("A", "B", "C"):
                raise ValueError("invalid_classification")
        if conversion == "storage":
            value = value.upper()
            if len(value) > 16:
                raise ValueError("invalid_storage_code")
        return value, state
    except (ValueError, DataError) as exc:
        code = "invalid_group_code" if isinstance(exc, DataError) else str(exc)
        parsed.issues.append(Problem(locator, code, identifier=pid, field=field, period=period))
        return None, "invalid"


def _source_id(cell: _Cell, parsed: Parsed, locator: str) -> tuple[str, str] | None:
    try:
        if cell.kind in ("f", "e", "d"):
            raise ValueError("invalid_identifier_cell")
        return str(cell.value), identifier(cell.value)
    except ValueError as exc:
        parsed.issues.append(Problem(locator, str(exc), field="source_identifier"))
        return None


def _quality(presence: dict[str, str]) -> str:
    return "invalid" if "invalid" in presence.values() else "valid"


def parse_cito(data: bytes, context: dict[str, Any]) -> Parsed:
    """Read the baseline Table XLSX or unambiguous six-field BIFF Cito schema."""
    current = _context_period(context, "current_period", required=True)
    target = _context_period(context, "target_period", required=True)
    if current == target:
        raise AdapterError("identical_cito_periods")
    if current is not None and target is not None and target < current:
        raise AdapterError("reversed_cito_periods")
    price_context = _price_context(context, cito=True)
    if (
        context.get("classification_context") not in ("ordinary", "dose")
        or not isinstance(context.get("group_namespace"), str)
        or not context["group_namespace"].strip()
    ):
        raise AdapterError("missing_classification_context")
    parsed = Parsed()
    try:
        sheet, mapping, rows = _cito_rows(data)
        for rowno, row in rows:
            if all(field_state(_at(row, col).value) == "blank" for col in mapping.values()):
                continue
            parsed.records += 1
            id_cell = _at(row, mapping["varenr"])
            ids = _source_id(id_cell, parsed, _locator(sheet, rowno, mapping["varenr"]))
            if ids is None:
                continue
            original, pid = ids
            raw_id = {
                "source_identifier": _raw(id_cell.value),
                "identifier_cell_type": id_cell.kind,
            }
            cost = _at(row, mapping["kostpris"])
            price_loc = _locator(sheet, rowno, mapping["kostpris"])
            amount, state = _field(
                cost, parsed, price_loc, pid, "amount_ore", current, conversion="money"
            )
            parsed.assertions.append(
                Assertion(
                    "price_observation",
                    price_loc,
                    original,
                    {**price_context, "amount_ore": amount},
                    {"amount_ore": state},
                    {**raw_id, "kostpris": _raw(cost.value), "kostpris_cell_type": cost.kind},
                    current,
                    current,
                    quality_state="invalid" if state == "invalid" else "valid",
                )
            )
            storage = _at(row, mapping["opbevkode"])
            storage_value, storage_state = _field(
                storage,
                parsed,
                _locator(sheet, rowno, mapping["opbevkode"]),
                pid,
                "storage_code",
                current,
                conversion="storage",
            )
            attributes = dict.fromkeys(_PACKAGE_FIELDS)
            attributes["storage_code"] = storage_value
            presence = dict.fromkeys(_PACKAGE_FIELDS, "unreported")
            presence["storage_code"] = storage_state
            parsed.assertions.append(
                Assertion(
                    "package_observation",
                    _locator(sheet, rowno),
                    original,
                    attributes,
                    presence,
                    {
                        **raw_id,
                        "opbevkode": _raw(storage.value),
                        "opbevkode_cell_type": storage.kind,
                    },
                    current,
                    current,
                    quality_state=_quality(presence),
                )
            )
            for heading, period in (("subkode", current), ("nysubkode", target)):
                cell = _at(row, mapping[heading])
                loc = _locator(sheet, rowno, mapping[heading])
                abc, abc_state = _field(
                    cell, parsed, loc, pid, "abc_code", period, conversion="abc"
                )
                values = {
                    "classification_context": context["classification_context"],
                    "abc_code": abc,
                    "group_namespace": context["group_namespace"],
                    "group_code": None,
                    "group_scope": None,
                }
                presence = {"abc_code": abc_state, "group_code": "unreported"}
                raw = {**raw_id, heading: _raw(cell.value)}
                if heading == "subkode":
                    group_cell = _at(row, mapping["subgrp"])
                    group, group_state = _field(
                        group_cell,
                        parsed,
                        _locator(sheet, rowno, mapping["subgrp"]),
                        pid,
                        "group_code",
                        current,
                        conversion="group",
                    )
                    values.update(
                        group_code=group, group_scope=current.isoformat() if current else None
                    )
                    presence["group_code"] = group_state
                    raw["subgrp"] = _raw(group_cell.value)
                    raw["subgrp_cell_type"] = group_cell.kind
                parsed.assertions.append(
                    Assertion(
                        "classification_observation",
                        loc,
                        original,
                        values,
                        presence,
                        raw,
                        period,
                        current,
                        quality_state=_quality(presence),
                    )
                )
        return parsed
    except (DataError, XLRDError, OSError, ValueError, IndexError) as exc:
        if isinstance(exc, AdapterError):
            raise
        raise AdapterError("invalid_cito_workbook") from exc


def parse_history(data: bytes, context: dict[str, Any]) -> Parsed:
    """Read exact-AIP rows and date columns without merging duplicate assertions."""
    current = _context_period(context, "current_period")
    price_context = _price_context(context, cito=False)
    parsed = Parsed()
    sheet = "Medicinpriser data"
    rows = _xlsx_rows(data, sheet)
    try:
        _, header = next(rows, (0, []))
        _validate_headers(header)
        headers = [c.value for c in header]
        mapping = _header_map(headers, ("varenummer", "indikator"))
        dates: dict[date, int] = {}
        for col, cell in enumerate(header):
            period = _period(cell.value)
            if period is not None:
                if period in dates:
                    raise AdapterError("duplicate_history_date")
                dates[period] = col
        if not dates:
            raise AdapterError("missing_history_dates")
        known = {"varenummer", "indikator", "indholdsstof", *_HISTORY_FIELDS}
        for col in range(min(dates.values()), max(dates.values()) + 1):
            if col not in dates.values() and normal_header(headers[col]) not in known:
                raise AdapterError("unknown_history_date_heading")
        attributes: dict[str, int] = {}
        for heading, field in _HISTORY_FIELDS.items():
            matches = [col for col, value in enumerate(headers) if normal_header(value) == heading]
            if len(matches) == 1:
                attributes[field] = matches[0]
            elif len(matches) > 1:
                parsed.issues.append(
                    Problem(
                        _locator(sheet, 1, matches[0]), "ambiguous_attribute_heading", field=field
                    )
                )
        for rowno, row in rows:
            indicator = _at(row, mapping["indikator"])
            if indicator.kind in ("f", "e"):
                parsed.issues.append(
                    Problem(
                        _locator(sheet, rowno, mapping["indikator"]),
                        "invalid_indicator_cell",
                        field="indicator",
                    )
                )
                continue
            if indicator.value != "AIP":
                continue
            parsed.records += 1
            id_cell = _at(row, mapping["varenummer"])
            ids = _source_id(id_cell, parsed, _locator(sheet, rowno, mapping["varenummer"]))
            if ids is None:
                continue
            original, pid = ids
            raw_id = {
                "source_identifier": _raw(id_cell.value),
                "identifier_cell_type": id_cell.kind,
            }
            for period, col in sorted(dates.items()):
                cell = _at(row, col)
                loc = _locator(sheet, rowno, col)
                amount, state = _field(
                    cell, parsed, loc, pid, "amount_ore", period, conversion="money"
                )
                parsed.assertions.append(
                    Assertion(
                        "price_observation",
                        loc,
                        original,
                        {**price_context, "amount_ore": amount},
                        {"amount_ore": state},
                        {**raw_id, "price": _raw(cell.value), "date_heading": _raw(headers[col])},
                        period,
                        current,
                        quality_state="invalid" if state == "invalid" else "valid",
                    )
                )
            values = dict.fromkeys(_PACKAGE_FIELDS)
            presence = dict.fromkeys(_PACKAGE_FIELDS, "unreported")
            raw = dict(raw_id)
            for field, col in attributes.items():
                cell = _at(row, col)
                values[field], presence[field] = _field(
                    cell,
                    parsed,
                    _locator(sheet, rowno, col),
                    pid,
                    field,
                    current,
                )
                raw[field] = _raw(cell.value)
            if attributes:
                loc = _locator(sheet, rowno)
                if current is None:
                    parsed.issues.append(
                        Problem(
                            loc,
                            "unknown_attribute_period",
                            identifier=pid,
                            field="effective_period",
                        )
                    )
                parsed.assertions.append(
                    Assertion(
                        "package_observation",
                        loc,
                        original,
                        values,
                        presence,
                        raw,
                        current,
                        current,
                        quality_state=_quality(presence),
                    )
                )
        return parsed
    except (DataError, XLRDError, OSError, ValueError, IndexError) as exc:
        if isinstance(exc, AdapterError):
            raise
        raise AdapterError("invalid_history_workbook") from exc
    finally:
        rows.close()
