"""Read-only, independent original-source reconciliation for the F2 report store.

This intentionally imports neither Store nor production adapters/converters. The
expected report contracts below are explicit; unsupported interpretations are
quarantined rather than counted as successful. Reports contain private evidence
and must be written outside the repository. Opening DuckDB is always read-only.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import struct
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path
from typing import Any, Never
from uuid import NAMESPACE_URL, UUID, uuid5

import duckdb
import olefile
import openpyxl
import xlrd
from openpyxl.utils import get_column_letter

from .archive import external_path

FIELDS = {
    "package_observation": (
        "name",
        "form",
        "strength",
        "pack_description",
        "company",
        "storage_code",
        "atc",
    ),
    "price_observation": (
        "price_type",
        "amount_ore",
        "currency",
        "price_basis",
        "tax_basis",
        "fee_basis",
        "price_role",
    ),
    "classification_observation": (
        "classification_context",
        "abc_code",
        "group_namespace",
        "group_code",
        "group_scope",
    ),
    "supply_observation": (
        "observation_date",
        "wholesaler_namespace",
        "wholesaler_identifier",
        "reported_status",
    ),
    "lifecycle_event": ("event_type", "original_source_code", "interpretation_status"),
}
HEADERS = {
    "price-changes": (
        "Firma",
        "Navn",
        "Form",
        "Styrke",
        "Vnr",
        "Pakningstekst",
        "Ny AIP",
        "Ny reg. pris",
        "Gl. AIP",
        "Gl. reg. pris",
        "Prisforskel %",
    ),
    "introductions": (
        "Navn",
        "Form",
        "Styrke",
        "MF-indehaver",
        "Pakningsejer",
        "Varenr",
        "Pakningsstørrelse",
        "Opbevaringsbetingelse",
        "Udlevering",
        "Trafikadvarsel DK",
        "AIP (Kr.)",
        "Reg. pris",
    ),
    "discontinued": (
        "Repræsentant",
        "Navn",
        "Form",
        "Styrke",
        "VareNr",
        "RPk",
        "Stop (Deaktiveret)",
        "Slet (Udgår permanent)",
    ),
    "supply": (
        "Varenr.",
        "Lægemiddel",
        "Form",
        "Styrke",
        "Pakningsstr.",
        "Firma",
        "Substitutionsgrp.nr.",
        "ABC-markering",
        "Grossist",
        "Ikrafttrædelsesdato",
    ),
}


def canonical(value: Any) -> str:
    return json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False
    )


def digest(value: Any) -> str:
    return hashlib.sha256(canonical(value).encode()).hexdigest()


def generated_id(kind: str, value: Any) -> str:
    return str(uuid5(NAMESPACE_URL, f"bef-redbird/foundation/v1/{kind}/{digest(value)}"))


def serial(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, (list, tuple)):
        return [serial(item) for item in value]
    if isinstance(value, dict):
        return {key: serial(item) for key, item in value.items()}
    return value


def persisted_record_hash(record: dict[str, Any], timestamp_columns: tuple[str, ...] = ()) -> str:
    """Match JSON default=str on SQL records, separately from ISO audit baselines.

    DATE and UUID already stringify identically. SQL datetimes use a space while
    source metadata/identity payloads intentionally use ISO timestamps with T.
    Only explicitly typed timestamp columns are converted; raw JSON is untouched.
    """
    converted = dict(record)
    for column in timestamp_columns:
        value = converted.get(column)
        if isinstance(value, str):
            converted[column] = str(datetime.fromisoformat(value))
    return digest(converted)


def rows(connection: duckdb.DuckDBPyConnection, table: str) -> list[dict[str, Any]]:
    if not re.fullmatch(r"[a-z_][a-z0-9_]*", table):
        raise ValueError("Invalid table")
    cur = connection.execute(f'SELECT * FROM "{table}"')
    names = [col[0] for col in cur.description]
    return [
        {key: serial(value) for key, value in zip(names, row, strict=True)}
        for row in cur.fetchall()
    ]


def snapshot(database: Path) -> dict[str, Any]:
    """Exact logical rows, including every original ID; never migrates the store."""
    result = {}
    with duckdb.connect(str(database), read_only=True) as con:
        # TIMESTAMPTZ display otherwise inherits the host session timezone. Event
        # instants and metadata derivations use UTC throughout the foundation.
        con.execute("SET TimeZone='UTC'")
        for (table,) in con.execute("SHOW TABLES").fetchall():
            records = sorted(rows(con, table), key=canonical)
            result[table] = {"count": len(records), "sha256": digest(records), "rows": records}
    return result


def snapshot_imports(database: Path, import_ids: list[str]) -> dict[str, Any]:
    """Read a bounded explicit import scope; never load unrelated frozen artifacts.

    Complete here means complete for these imports. The corpus inventory must
    independently cover every import when assembling whole-archive acceptance.
    """
    selected = sorted({str(UUID(value)) for value in import_ids})
    if not selected or len(selected) != len(import_ids):
        raise ValueError("Explicit audit scope requires unique existing imports")
    result: dict[str, Any] = {}
    with duckdb.connect(str(database), read_only=True) as con:
        con.execute("SET TimeZone='UTC'")
        marks = ",".join("?" for _ in selected)
        filters = {
            "import_batch": f"import_id IN ({marks})",
            "import_issue": f"import_id IN ({marks})",
            "price_period": (
                "period_start IN (SELECT current_period FROM source_delivery "
                "WHERE delivery_id IN (SELECT delivery_id FROM import_batch "
                f"WHERE import_id IN ({marks})))"
            ),
            "source_delivery": (
                "delivery_id IN (SELECT delivery_id FROM import_batch "
                f"WHERE import_id IN ({marks}))"
            ),
            "package_identifier": f"evidence_import_id IN ({marks})",
            "package": (
                "package_id IN (SELECT package_id FROM package_identifier "
                f"WHERE evidence_import_id IN ({marks}))"
            ),
            **{table: f"import_id IN ({marks})" for table in FIELDS},
        }
        for table, where in filters.items():
            cur = con.execute(f'SELECT * FROM "{table}" WHERE {where}', selected)
            names = [col[0] for col in cur.description]
            records = sorted(
                [
                    {key: serial(value) for key, value in zip(names, row, strict=True)}
                    for row in cur.fetchall()
                ],
                key=canonical,
            )
            result[table] = {"count": len(records), "sha256": digest(records), "rows": records}
        dependencies = _basis_dependency_ids(result["import_batch"]["rows"])
        known = {row["delivery_id"] for row in result["source_delivery"]["rows"]}
        extra = sorted(dependencies - known)
        if extra:
            placeholders = ",".join("?" for _ in extra)
            cur = con.execute(
                f"SELECT * FROM source_delivery WHERE delivery_id IN ({placeholders})", extra
            )
            names = [col[0] for col in cur.description]
            records = result["source_delivery"]["rows"] + [
                {key: serial(value) for key, value in zip(names, row, strict=True)}
                for row in cur.fetchall()
            ]
            records.sort(key=canonical)
            result["source_delivery"] = {
                "count": len(records),
                "sha256": digest(records),
                "rows": records,
            }
        if result["import_batch"]["count"] != len(selected):
            raise ValueError("Explicit audit scope contains missing imports")
    return result


def compare_baseline(database: Path, baseline: dict[str, Any]) -> dict[str, Any]:
    """Require every protected row exactly without loading the later archive."""
    result = {}
    with duckdb.connect(str(database), read_only=True) as con:
        con.execute("SET TimeZone='UTC'")
        for table, previous in baseline.items():
            if not re.fullmatch(r"[a-z_][a-z0-9_]*", table):
                raise ValueError("Invalid baseline table")
            keys = con.execute("SELECT name FROM pragma_table_info(?) WHERE pk", [table]).fetchall()
            key_names = [key[0] for key in keys]
            if not keys or any(not re.fullmatch(r"[a-z_][a-z0-9_]*", key) for key in key_names):
                raise ValueError("Baseline comparison requires a supported primary key")
            old = Counter(canonical(row) for row in previous["rows"])
            new: Counter[str] = Counter()
            identifiers = [[row[key] for key in key_names] for row in previous["rows"]]
            for offset in range(0, len(identifiers), 512):
                values = identifiers[offset : offset + 512]
                if len(key_names) == 1:
                    marks = ",".join("?" for _ in values)
                    condition = f'"{key_names[0]}" IN ({marks})'
                else:
                    term = " AND ".join(f'"{key}"=?' for key in key_names)
                    condition = " OR ".join(f"({term})" for _ in values)
                cur = con.execute(
                    f'SELECT * FROM "{table}" WHERE {condition}',
                    [value for identifier in values for value in identifier],
                )
                columns = [column[0] for column in cur.description]
                new.update(
                    canonical({k: serial(v) for k, v in zip(columns, row, strict=True)})
                    for row in cur.fetchall()
                )
            total = con.execute(f'SELECT count(*) FROM "{table}"').fetchone()
            assert total is not None
            preserved = sum((old & new).values())
            result[table] = {
                "missing_or_changed": sum(old.values()) - preserved,
                "additional": total[0] - preserved,
            }
    return {
        "preserved": all(x["missing_or_changed"] == 0 for x in result.values()),
        "tables": result,
    }


def _blank(value: Any) -> bool:
    return value is None or isinstance(value, str) and not value.strip()


def _raw(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else value


def _money(value: Any) -> int:
    if isinstance(value, bool):
        raise ValueError("malformed_money")
    if isinstance(value, str):
        token = value.strip()
        if not re.fullmatch(r"-?(?:\d+|\d{1,3}(?:\.\d{3})+)(?:,\d+)?", token, re.ASCII):
            raise ValueError("malformed_money")
        token = token.replace(".", "").replace(",", ".")
    else:
        token = str(value)
    try:
        number = Decimal(token)
    except InvalidOperation as exc:
        raise ValueError("malformed_money") from exc
    if not number.is_finite():
        raise ValueError("nonfinite_money")
    if number < 0:
        raise ValueError("negative_money")
    with localcontext() as ctx:
        ctx.prec = max(28, len(number.as_tuple().digits) + 3)
        ore = number * 100
    if ore != ore.to_integral_value():
        raise ValueError("unsupported_money_precision")
    if ore > 2**63 - 1:
        raise ValueError("money_overflow")
    return int(ore)


def _date(value: Any) -> str:
    if isinstance(value, datetime):
        if value.time() != datetime.min.time():
            raise ValueError("invalid_date")
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d-%m-%Y"):
            try:
                converted = datetime.strptime(value.strip(), fmt)
                if converted.strftime(fmt) == value.strip():
                    return converted.date().isoformat()
            except ValueError:
                pass
    raise ValueError("invalid_date")


def _biff_formulas(data: bytes) -> dict[int, set[tuple[int, int]]]:
    """Independent BIFF record scan, with sheet starts from BOUNDSHEET records."""
    with olefile.OleFileIO(data) as container:
        stream_name = "Workbook" if container.exists("Workbook") else "Book"  # type: ignore[arg-type]
        stream = container.openstream(stream_name).read()  # type: ignore[arg-type]
    bounds = []
    formulas = []
    offset = 0
    while offset + 4 <= len(stream):
        tag, length = struct.unpack_from("<HH", stream, offset)
        payload = stream[offset + 4 : offset + 4 + length]
        if tag == 0x0085 and len(payload) >= 4:
            bounds.append(struct.unpack_from("<I", payload)[0])
        if tag == 0x0006 and len(payload) >= 4:
            formulas.append((offset, *struct.unpack_from("<HH", payload)))
        offset += 4 + length
    return {
        i: {
            (row, col)
            for pos, row, col in formulas
            if start <= pos < min([b for b in bounds if b > start] or [len(stream)])
        }
        for i, start in enumerate(bounds)
    }


def read_grids(data: bytes) -> list[dict[str, Any]]:
    """Read actual cells using raw openpyxl/xlrd, preserving serial and kind."""
    grids: list[dict[str, Any]] = []
    if data.startswith(b"PK"):
        book = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=True)
        try:
            for sheet in book:
                # Some official exports understate their XML dimension and would
                # otherwise hide actual cells, including lifecycle date columns.
                sheet.reset_dimensions()
                grids.append(
                    {
                        "name": sheet.title,
                        "date_system": None,
                        "rows": [
                            [(c.value, c.data_type, _raw(c.value)) for c in row]
                            for row in sheet.iter_rows()
                        ],
                    }
                )
        finally:
            book.close()
    elif data.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
        formulas = _biff_formulas(data)
        legacy = xlrd.open_workbook(file_contents=data, logfile=io.StringIO())
        try:
            for index, legacy_sheet in enumerate(legacy.sheets()):
                grid_rows = []
                for r in range(legacy_sheet.nrows):
                    record = []
                    for c in range(legacy_sheet.ncols):
                        legacy_cell = legacy_sheet.cell(r, c)
                        value: Any = legacy_cell.value
                        kind = "n"
                        if (r, c) in formulas.get(index, set()):
                            kind = "f"
                        elif legacy_cell.ctype == xlrd.XL_CELL_ERROR:
                            kind = "e"
                        elif legacy_cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value, kind = bool(value), "b"
                        elif legacy_cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(
                                float(legacy_cell.value), legacy.datemode
                            )
                        record.append((value, kind, _raw(legacy_cell.value)))
                    grid_rows.append(record)
                grids.append(
                    {"name": legacy_sheet.name, "date_system": legacy.datemode, "rows": grid_rows}
                )
        finally:
            legacy.release_resources()
    else:
        raise ValueError("unsupported_source_format")
    return grids


def _basis_override(context: dict[str, Any], *, enabled: bool) -> dict[str, Any] | None:
    """Independently validate declared scope syntax, not the evidence's meaning."""
    supplied = "price_basis_override" in context
    if not enabled:
        if supplied:
            raise ValueError("independent_price_basis_override_version_unavailable")
        return None
    profile = context.get("price_basis_override")
    expected_keys = {
        "version",
        "price_type",
        "price_role",
        "currency",
        "price_basis",
        "tax_basis",
        "fee_basis",
        "effective_period_start_from",
        "effective_period_start_through",
        "evidence_delivery_id",
        "evidence_sha256",
        "evidence_locator",
    }
    if not isinstance(profile, dict) or set(profile) != expected_keys:
        raise ValueError("independent_invalid_price_basis_override")
    if type(profile["version"]) is not int or profile["version"] != 1:
        raise ValueError("independent_invalid_price_basis_override")
    actual_scope = tuple(
        profile[k]
        for k in ("price_type", "price_role", "currency", "price_basis", "tax_basis", "fee_basis")
    )
    if actual_scope != (
        "aip",
        "effective_price",
        "DKK",
        "per_package",
        "excluded",
        "excluded",
    ) or any(context.get(k) != "unknown" for k in ("tax_basis", "fee_basis")):
        raise ValueError("independent_invalid_price_basis_override")
    try:
        left, right = (
            profile["effective_period_start_from"],
            profile["effective_period_start_through"],
        )
        if not isinstance(left, str) or not isinstance(right, str):
            raise ValueError
        if (
            date.fromisoformat(left).isoformat() != left
            or date.fromisoformat(right).isoformat() != right
            or left > right
        ):
            raise ValueError
        evidence_id = profile["evidence_delivery_id"]
        fingerprint = profile["evidence_sha256"]
        locator = profile["evidence_locator"]
        if not isinstance(evidence_id, str) or str(UUID(evidence_id)) != evidence_id:
            raise ValueError
        if not isinstance(fingerprint, str) or re.fullmatch("[a-f0-9]{64}", fingerprint) is None:
            raise ValueError
        if not isinstance(locator, str) or not locator.strip():
            raise ValueError
    except (ValueError, TypeError, AttributeError) as exc:
        raise ValueError("independent_invalid_price_basis_override") from exc
    return dict(profile)


def _basis_dependency_ids(batches: list[dict[str, Any]]) -> set[str]:
    result = set()
    for batch in batches:
        try:
            adapter_context = json.loads(batch["interpretation_context"]).get("adapter_context", {})
            profile = _basis_override(
                adapter_context,
                enabled=batch["parser_version"]
                in {"price-changes-workbook-v2", "introductions-workbook-v3", "lms-ftp-zip-v1"},
            )
            if profile is not None:
                result.add(profile["evidence_delivery_id"])
        except (ValueError, KeyError, TypeError, AttributeError):
            # Invalid declarations are quarantined by the import audit, not hidden.
            continue
    return result


def _basis_evidence_pin(
    profile: dict[str, Any], deliveries: dict[str, Any], archive: Path
) -> dict[str, Any]:
    """Verify immutable referenced bytes; semantic applicability needs external review."""
    evidence_id = profile["evidence_delivery_id"]
    delivery = deliveries.get(evidence_id)
    if delivery is None or delivery["content_hash"] != profile["evidence_sha256"]:
        raise ValueError("independent_price_basis_evidence_pin_mismatch")
    path = (archive / delivery["archive_location"]).resolve()
    if (
        not path.is_relative_to(archive.resolve())
        or not path.is_file()
        or hashlib.sha256(path.read_bytes()).hexdigest() != profile["evidence_sha256"]
    ):
        raise ValueError("independent_price_basis_evidence_archive_mismatch")
    return {
        "delivery_id": evidence_id,
        "content_hash": profile["evidence_sha256"],
        "source_locator": profile["evidence_locator"],
        "validation": (
            "Referenced archive byte pin only; meaning/applicability require separately "
            "reviewed evidence and availability remains selection-gated"
        ),
    }


def expected_report(
    data: bytes,
    adapter: str,
    context: dict[str, Any],
    *,
    parser_version: str | None = None,
) -> dict[str, Any]:
    """Independent expected assertions for the four report families in the F2 store.

    Bounded history v2 has a separate streaming contract. Cito, full history v1
    and optional supply classification never pass without independent contracts.
    """
    version = parser_version or (
        "lms-history-workbook-v2"
        if adapter == "lms-history" and "subset_identifiers" in context
        else adapter + "-workbook-v1"
    )
    allowed_versions = {adapter + "-workbook-v1"}
    if adapter in ("introductions", "discontinued", "lms-history"):
        allowed_versions.add(adapter + "-workbook-v2")
    if adapter == "price-changes":
        allowed_versions.add("price-changes-workbook-v2")
    if adapter == "introductions":
        allowed_versions.add("introductions-workbook-v3")
    if version not in allowed_versions:
        raise ValueError("independent_version_contract_unavailable")
    if adapter == "lms-history" and "subset_identifiers" in context:
        if version != "lms-history-workbook-v2":
            raise ValueError("independent_version_contract_unavailable")
        return expected_history_subset(data, context)
    if adapter not in HEADERS or adapter == "supply" and context.get("classification_context"):
        raise ValueError("independent_contract_unavailable")
    basis_override = _basis_override(
        context, enabled=version in {"price-changes-workbook-v2", "introductions-workbook-v3"}
    )
    expanded = version in (
        "introductions-workbook-v2",
        "introductions-workbook-v3",
        "discontinued-workbook-v2",
    )
    if expanded and not data.startswith(b"PK"):
        raise ValueError("independent_unsupported_source_format")
    candidates = []
    for grid in read_grids(data):
        for h in {
            "price-changes": (4,),
            "introductions": (5, 6),
            "discontinued": (1, 2) if expanded else (1,),
            "supply": (1,),
        }[adapter]:
            if h > len(grid["rows"]):
                continue
            headings = [
                " ".join(str(c[0] or "").replace("_x000D_", " ").split())
                for c in grid["rows"][h - 1]
            ]
            accepted = list(HEADERS[adapter])
            if adapter == "discontinued" and h == 2:
                accepted = [
                    "Repræsentant",
                    "Lægemiddel",
                    "Lægemiddelform",
                    "Styrke",
                    "Varenummer",
                    "RPk",
                    "Stop (deaktiveret)",
                    "Slet (udgår permanent)",
                ]
            if adapter == "introductions" and headings[5:6] == ["Varenr."]:
                accepted[5] = "Varenr."
            if (
                headings[: len(accepted)] == accepted
                and not any(headings[len(accepted) :])
                and (not expanded or not any(c[1] in ("f", "e") for c in grid["rows"][h - 1]))
            ):
                candidates.append((grid, h, headings))
    if len(candidates) != 1:
        raise ValueError("independent_schema_mismatch")
    grid, header, headings = candidates[0]
    sheet = grid["name"]
    current = context.get("current_period")
    issues: list[dict[str, Any]] = []
    assertions: list[dict[str, Any]] = []
    source_rows = 0

    def loc(r: int, c: int | None = None) -> str:
        return f"sheet:{sheet}!" + (f"row:{r}" if c is None else f"{get_column_letter(c + 1)}{r}")

    def issue(
        r: int,
        c: int | None,
        code: str,
        identifier: str | None,
        field: str,
        period: str | None = None,
        warning: bool = False,
    ) -> None:
        issues.append(
            {
                "source_locator": loc(r, c),
                "code": code,
                "identifier": identifier,
                "field": field,
                "period": period,
                "severity": "warning" if warning else "error",
                "blocking_scope": "none" if warning else "field",
            }
        )

    metadata: dict[str, Any] = {}
    if adapter == "discontinued" and header == 2:
        if current is not None:
            raise ValueError("independent_unestablished_report_period")
        first = grid["rows"][0]
        if (
            not first
            or first[0][1] in ("f", "e")
            or not isinstance(first[0][0], (date, datetime))
            or any(not _blank(c[0]) for c in first[1:])
        ):
            raise ValueError("independent_metadata_mismatch")
        _date(first[0][0])  # Reject times; retain the exact typed source representation.
        metadata.update(
            unlabelled_date=_raw(first[0][0]),
            unlabelled_date_locator=loc(1, 0),
            unlabelled_date_semantics="unknown",
        )
    if adapter in ("price-changes", "introductions"):
        title = grid["rows"][0][0][0]
        allowed = (
            ["Prisændringer"]
            if adapter == "price-changes"
            else [
                "Nye pakninger på eksisterende formuleringer",
                "Nye pakninger på nye formuleringer",
            ]
        )
        if title not in allowed or expanded and grid["rows"][0][0][1] in ("f", "e"):
            raise ValueError("independent_title_mismatch")
        metadata["report_title"] = title
        locations = (
            [(1, 4), (1, 5), (1, 11)]
            if adapter == "price-changes"
            else [(2, 6), (2, 7), (2, 8), (2, 9), (3, 12)]
        )
        if expanded and adapter == "introductions":
            locations = [(2, c) for c in range(4, 12)] + [(3, 12), (3, 13), (4, 12)]
            for r, row in enumerate(grid["rows"][: header - 1], 1):
                for c, cell in enumerate(row):
                    if (r, c) not in locations and isinstance(cell[0], str):
                        if re.match(r"(?:Takstdato|Kørselsdato)\s*:", cell[0]):
                            raise ValueError("independent_unrecognized_metadata_location")
        for r, c in locations:
            row = grid["rows"][r - 1] if r <= len(grid["rows"]) else []
            value = row[c][0] if c < len(row) else None
            if _blank(value):
                continue
            if expanded and (row[c][1] in ("f", "e") or not isinstance(value, str)):
                raise ValueError("independent_metadata_mismatch")
            match = re.fullmatch(r"(Takstdato|Kørselsdato):\s*(\d{2}-\d{2}-\d{4})", str(value))
            if match is None:
                raise ValueError("independent_metadata_mismatch")
            if expanded and match[1] != ("Takstdato" if r == 2 else "Kørselsdato"):
                raise ValueError("independent_metadata_mismatch")
            key = "tariff_date" if match[1] == "Takstdato" else "run_date"
            observed = _date(match[2])
            if key in metadata and metadata[key] != observed:
                raise ValueError("independent_context_conflict")
            metadata.update({key: observed, key + "_source": value, key + "_locator": loc(r, c)})
        if metadata.get("tariff_date") != current:
            raise ValueError("independent_context_conflict")
        if adapter == "price-changes" and "run_date" not in metadata:
            raise ValueError("independent_run_date_missing")

    for r, cells in enumerate(grid["rows"][header:], header + 1):
        if all(_blank(c[0]) for c in cells):
            continue
        source_rows += 1

        def cell(c: int, source_cells: list[Any] = cells) -> tuple[Any, str, Any]:
            return source_cells[c] if c < len(source_cells) else (None, "n", None)

        id_col = {"supply": 0, "introductions": 5}.get(adapter, 4)
        original = str(cell(id_col)[0])
        if cell(id_col)[1] in ("f", "e", "d") or not re.fullmatch(
            r"[0-9]{1,6}(?:\.0+)?", original.strip()
        ):
            issue(
                r,
                id_col,
                "invalid_identifier_cell"
                if cell(id_col)[1] in ("f", "e", "d")
                else "invalid_identifier",
                None,
                "source_identifier",
            )
            continue
        canonical_id = original.strip().split(".")[0].zfill(6)
        raw = {h: cell(c)[2] for c, h in enumerate(headings) if h}
        raw.update(source_identifier=_raw(cell(id_col)[0]), identifier_cell_type=cell(id_col)[1])
        raw["source_cell_types"] = {h: cell(c)[1] for c, h in enumerate(headings) if h}
        raw["source_field_presence"] = {
            h: "invalid"
            if cell(c)[1] in ("f", "e")
            else "blank"
            if _blank(cell(c)[0])
            else "present"
            for c, h in enumerate(headings)
            if h
        }
        if grid["date_system"] is not None:
            raw["excel_date_system"] = grid["date_system"]
        raw.update(metadata)

        def field(
            c: int,
            name: str,
            conversion: str = "text",
            period: str | None = current,
            row_number: int = r,
            row_identifier: str = canonical_id,
        ) -> tuple[Any, str]:
            value, kind, _ = cell(c)
            try:
                if kind in ("f", "e"):
                    raise ValueError("formula_cell" if kind == "f" else "error_cell")
                if _blank(value):
                    return None, "blank"
                converted = (
                    _money(value)
                    if conversion == "money"
                    else _date(value)
                    if (conversion == "date")
                    else str(value).strip()
                )
                return converted, "present"
            except ValueError as exc:
                issue(row_number, c, str(exc), row_identifier, name, period)
                return None, "invalid"

        def assertion(
            table: str,
            c: int | None,
            values: dict[str, Any],
            presence: dict[str, str],
            effective_period: str | None = None,
            effective_date: str | None = None,
            quality: str = "valid",
            extra_raw: dict[str, Any] | None = None,
            row_number: int = r,
            row_identifier: str = canonical_id,
            original_identifier: str = original,
            raw_evidence: dict[str, Any] = raw,
        ) -> None:
            assertions.append(
                {
                    "table": table,
                    "source_locator": loc(row_number, c),
                    "source_identifier": original_identifier,
                    "canonical_identifier": row_identifier,
                    "effective_period": effective_period,
                    "report_period": current,
                    "effective_date": effective_date,
                    "values": values,
                    "presence": presence,
                    "quality": quality,
                    "raw": {**raw_evidence, **(extra_raw or {})},
                }
            )

        attrs = dict.fromkeys(FIELDS["package_observation"])
        presence = dict.fromkeys(attrs, "unreported")
        columns = {"company": 0, "name": 1, "form": 2, "strength": 3, "pack_description": 5}
        if adapter == "introductions":
            columns = {
                "name": 0,
                "form": 1,
                "strength": 2,
                "company": 4,
                "pack_description": 6,
                "storage_code": 7,
            }
        if adapter == "supply":
            columns = {"name": 1, "form": 2, "strength": 3, "pack_description": 4, "company": 5}
        for name, c in columns.items():
            attrs[name], presence[name] = field(c, name)
        assertion("package_observation", None, attrs, presence, current)
        if adapter in ("introductions", "price-changes"):
            prices = [(10, "aip", "effective_price"), (11, "registered", "effective_price")]
            if adapter == "price-changes":
                prices = [
                    (6, "aip", "effective_price"),
                    (7, "registered", "effective_price"),
                    (8, "aip", "reported_previous_price"),
                    (9, "registered", "reported_previous_price"),
                ]
            for c, price_type, role in prices:
                period = current if role == "effective_price" else None
                amount, state = field(c, "amount_ore", "money", period)
                tax_basis, fee_basis = context["tax_basis"], context["fee_basis"]
                if (
                    basis_override is not None
                    and price_type == "aip"
                    and role == "effective_price"
                    and period is not None
                    and basis_override["effective_period_start_from"]
                    <= period
                    <= basis_override["effective_period_start_through"]
                ):
                    tax_basis = fee_basis = "excluded"
                assertion(
                    "price_observation",
                    c,
                    {
                        "price_type": price_type,
                        "amount_ore": amount,
                        "currency": "DKK",
                        "price_basis": "per_package",
                        "tax_basis": tax_basis,
                        "fee_basis": fee_basis,
                        "price_role": role,
                    },
                    {"amount_ore": state},
                    period,
                )
            if adapter == "price-changes":
                field(10, "reported_registered_percentage_change")
            else:
                assertion(
                    "lifecycle_event",
                    None,
                    {
                        "event_type": "introduction",
                        "original_source_code": metadata["report_title"],
                        "interpretation_status": "established",
                    },
                    {"effective_date": "present"},
                    effective_date=current,
                )
        if adapter == "discontinued":
            for c, event in [(6, "deactivation"), (7, "permanent_deletion")]:
                effective, state = field(c, "effective_date", "date", None)
                if effective is None:
                    if state == "blank":
                        issue(
                            r, c, "missing_event_date", canonical_id, "effective_date", warning=True
                        )
                    continue
                issue(
                    r,
                    c,
                    "source_reported_event_date_semantics_unknown",
                    canonical_id,
                    "effective_date",
                    warning=True,
                )
                assertion(
                    "lifecycle_event",
                    c,
                    {
                        "event_type": event,
                        "original_source_code": grid["rows"][header - 1][c][0],
                        "interpretation_status": "unknown",
                    },
                    {"effective_date": state},
                    effective_date=effective,
                    quality="incomplete",
                )
        if adapter == "supply":
            wholesaler, ws = field(8, "wholesaler_identifier")
            effective, ds = field(9, "effective_date", "date", None)
            source_date = context.get("date_basis") == "source_effective_date"
            row_observed = effective if source_date else context.get("observation_date")
            if wholesaler is None:
                issue(r, 8, "missing_wholesaler_identifier", canonical_id, "wholesaler_identifier")
            if row_observed is None:
                issue(r, 9, "missing_supply_observation_date", canonical_id, "observation_date")
            if row_observed is not None and wholesaler is not None:
                role = "source_effective_date" if source_date else "declared_observation_date"
                assertion(
                    "supply_observation",
                    8,
                    {
                        "observation_date": row_observed,
                        "wholesaler_namespace": context["wholesaler_namespace"],
                        "wholesaler_identifier": wholesaler,
                        "reported_status": context["reported_status"],
                    },
                    {
                        "wholesaler_identifier": ws,
                        "effective_date": ds,
                        "observation_date": "present",
                        "reported_status": "present",
                    },
                    effective_date=effective,
                    extra_raw={
                        "observation_date_role": role,
                        "observation_date_basis": "source_effective_date"
                        if source_date
                        else context["observation_date_basis"],
                    },
                )
    return {
        "assertions": assertions,
        "issues": issues,
        "source_rows": source_rows,
        "sheet": sheet,
        "header_row": header,
        "metadata": metadata,
        "date_system": grid["date_system"],
    }


def expected_history_subset(data: bytes, context: dict[str, Any]) -> dict[str, Any]:
    """Independent streaming history expectations; scope before derivation.

    All AIP rows and unselected fields remain accounted as explicitly out of scope.
    Dates are read from header cells, never from filenames or workbook properties.
    """
    wanted_ids = context.get("subset_identifiers")
    wanted_periods = context.get("subset_periods")
    if (
        not isinstance(wanted_ids, list)
        or not 1 <= len(wanted_ids) <= 5
        or any(not isinstance(v, str) or not re.fullmatch(r"[0-9]{6}", v) for v in wanted_ids)
        or len(set(wanted_ids)) != len(wanted_ids)
    ):
        raise ValueError("independent_invalid_history_subset_identifiers")
    if (
        not isinstance(wanted_periods, list)
        or not 1 <= len(wanted_periods) <= 2
        or any(not isinstance(p, str) for p in wanted_periods)
        or len(set(wanted_periods)) != len(wanted_periods)
        or any(_date(p) != p for p in wanted_periods)
    ):
        raise ValueError("independent_invalid_history_subset_periods")
    current = context.get("current_period")
    sheet_name = "Medicinpriser data"
    book = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    result: dict[str, Any] = {
        "assertions": [],
        "issues": [],
        "source_rows": 0,
        "sheet": sheet_name,
        "header_row": 1,
        "metadata": {},
        "date_system": None,
        "scope_accounting": {
            "physical_data_rows": 0,
            "aip_rows": 0,
            "selected_aip_rows": 0,
            "unselected_aip_rows": 0,
            "non_aip_rows": 0,
            "invalid_unscoped_identifiers": 0,
            "selected_invalid_indicators": 0,
            "selected_price_cells": 0,
            "unselected_date_fields_per_selected_row": 0,
        },
    }
    accounting = result["scope_accounting"]
    try:
        sheet = book[sheet_name]
        it = sheet.iter_rows()
        header = next(it)
        headings = [" ".join(str(c.value or "").lower().split()) for c in header]
        if any(c.data_type in ("f", "e") for c in header):
            raise ValueError("independent_invalid_header")
        if headings.count("varenummer") != 1 or headings.count("indikator") != 1:
            raise ValueError("independent_history_schema_mismatch")
        id_col, indicator_col = headings.index("varenummer"), headings.index("indikator")
        dates = {}
        for col, cell in enumerate(header):
            value = cell.value
            period = None
            if isinstance(value, (date, datetime)):
                period = (
                    value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
                )
            elif re.fullmatch(r"[0-9]{8}", str(value)):
                period = datetime.strptime(str(value), "%Y%m%d").date().isoformat()
            if period is not None:
                if period in dates:
                    raise ValueError("independent_duplicate_history_date")
                dates[period] = col
        if not set(wanted_periods) <= dates.keys():
            raise ValueError("independent_missing_subset_date")
        accounting["unselected_date_fields_per_selected_row"] = len(dates) - len(wanted_periods)
        attribute_headers = {
            "lægemiddel": "name",
            "form": "form",
            "styrke": "strength",
            "pakning": "pack_description",
            "firma": "company",
            "atc": "atc",
        }
        attributes = {}
        for heading, field in attribute_headers.items():
            matches = [i for i, h in enumerate(headings) if h == heading]
            if len(matches) == 1:
                attributes[field] = matches[0]
            elif len(matches) > 1:
                result["issues"].append(
                    {
                        "source_locator": (
                            f"sheet:{sheet_name}!{get_column_letter(matches[0] + 1)}1"
                        ),
                        "code": "ambiguous_attribute_heading",
                        "identifier": None,
                        "field": field,
                        "period": None,
                        "severity": "error",
                        "blocking_scope": "field",
                    }
                )
        result["metadata"] = {
            "subset_identifiers": wanted_ids,
            "subset_periods": wanted_periods,
            "date_header_evidence": {
                p: {
                    "locator": f"sheet:{sheet_name}!{get_column_letter(dates[p] + 1)}1",
                    "value": _raw(header[dates[p]].value),
                }
                for p in wanted_periods
            },
        }
        for rn, row in enumerate(it, 2):
            accounting["physical_data_rows"] += 1
            value = row[id_col].value
            token = str(value).strip()
            canonical_id = (
                token.split(".")[0].zfill(6)
                if not isinstance(value, bool) and re.fullmatch(r"[0-9]{1,6}(?:\.0+)?", token)
                else None
            )
            indicator = row[indicator_col]
            if indicator.value == "AIP":
                accounting["aip_rows"] += 1
                if canonical_id not in wanted_ids:
                    accounting["unselected_aip_rows"] += 1
                    if canonical_id is None:
                        accounting["invalid_unscoped_identifiers"] += 1
            else:
                accounting["non_aip_rows"] += 1
            if canonical_id not in wanted_ids:
                continue

            def locator(col: int | None, number: int = rn) -> str:
                return f"sheet:{sheet_name}!" + (
                    f"row:{number}" if col is None else f"{get_column_letter(col + 1)}{number}"
                )

            def issue(
                col: int | None,
                code: str,
                field: str,
                period: str | None = None,
                identifier: str | None = canonical_id,
            ) -> None:
                result["issues"].append(
                    {
                        "source_locator": locator(col),
                        "code": code,
                        "identifier": identifier,
                        "field": field,
                        "period": period,
                        "severity": "error",
                        "blocking_scope": "field",
                    }
                )

            if indicator.data_type in ("f", "e"):
                accounting["selected_invalid_indicators"] += 1
                issue(indicator_col, "invalid_indicator_cell", "indicator", identifier=None)
                continue
            if indicator.value != "AIP":
                continue
            accounting["selected_aip_rows"] += 1
            result["source_rows"] += 1
            id_cell = row[id_col]
            if id_cell.data_type in ("f", "e", "d"):
                issue(id_col, "invalid_identifier_cell", "source_identifier", identifier=None)
                continue
            raw_id = {"source_identifier": _raw(value), "identifier_cell_type": id_cell.data_type}

            def history_field(
                col: int,
                name: str,
                period: str | None,
                money: bool = False,
                source_row: Any = row,
            ) -> tuple[Any, str]:
                cell = source_row[col]
                try:
                    if cell.data_type in ("f", "e", "d"):
                        raise ValueError(
                            {"f": "formula_cell", "e": "error_cell", "d": "invalid_date_cell"}[
                                cell.data_type
                            ]
                        )
                    if _blank(cell.value):
                        return None, "blank"
                    return (_money(cell.value) if money else str(cell.value).strip()), "present"
                except ValueError as exc:
                    issue(col, str(exc), name, period)
                    return None, "invalid"

            for period in sorted(wanted_periods):
                col = dates[period]
                amount, state = history_field(col, "amount_ore", period, True)
                accounting["selected_price_cells"] += 1
                result["assertions"].append(
                    {
                        "table": "price_observation",
                        "source_locator": locator(col),
                        "source_identifier": str(value),
                        "canonical_identifier": canonical_id,
                        "effective_period": period,
                        "report_period": current,
                        "effective_date": None,
                        "values": {
                            "price_type": "aip",
                            "currency": "DKK",
                            "price_basis": "per_package",
                            "tax_basis": context["tax_basis"],
                            "fee_basis": context["fee_basis"],
                            "price_role": "effective_price",
                            "amount_ore": amount,
                        },
                        "presence": {"amount_ore": state},
                        "quality": "invalid" if state == "invalid" else "valid",
                        "raw": {
                            **raw_id,
                            "price": _raw(row[col].value),
                            "date_heading": _raw(header[col].value),
                        },
                    }
                )
            if attributes:
                values = dict.fromkeys(FIELDS["package_observation"])
                presence = dict.fromkeys(values, "unreported")
                raw = dict(raw_id)
                for name, col in attributes.items():
                    values[name], presence[name] = history_field(col, name, current)
                    raw[name] = _raw(row[col].value)
                if current is None:
                    issue(None, "unknown_attribute_period", "effective_period")
                result["assertions"].append(
                    {
                        "table": "package_observation",
                        "source_locator": locator(None),
                        "source_identifier": str(value),
                        "canonical_identifier": canonical_id,
                        "effective_period": current,
                        "report_period": current,
                        "effective_date": None,
                        "values": values,
                        "presence": presence,
                        "raw": raw,
                        "quality": "invalid" if "invalid" in presence.values() else "valid",
                    }
                )
        return result
    finally:
        book.close()


def auditor_fingerprint() -> dict[str, str]:
    """Pin all independent readers, including format-specific source accounting."""
    root = Path(__file__).parent
    return {
        name: hashlib.sha256((root / name).read_bytes()).hexdigest()
        for name in ("source_audit.py", "rtf_audit.py", "pdf_audit.py", "lms_ftp_audit.py")
        if (root / name).is_file()
    }


def reconcile(
    database: Path,
    archive: Path,
    original_paths: list[Path] | None = None,
    *,
    import_ids: list[str] | None = None,
) -> dict[str, Any]:
    """All existing records accounted; results and mismatches are private output."""
    baseline = snapshot(database) if import_ids is None else snapshot_imports(database, import_ids)
    tables = {table: content["rows"] for table, content in baseline.items()}
    report: dict[str, Any] = {
        "version": 1,
        "auditor_code": auditor_fingerprint(),
        "audit_scope": (
            {"kind": "whole_database"}
            if import_ids is None
            else {"kind": "explicit_imports", "import_ids": sorted(import_ids)}
        ),
        "audited_database_sha256": (
            hashlib.sha256(database.read_bytes()).hexdigest() if import_ids is None else None
        ),
        "database_fingerprint_basis": (
            "closed_database_bytes" if import_ids is None else "scoped_persisted_record_hashes"
        ),
        "baseline_tables": {
            t: {"count": v["count"], "sha256": v["sha256"]} for t, v in baseline.items()
        },
        "imports": [],
        "mismatches": [],
        "verified_observation_ids": [],
        "quarantined_observation_ids": [],
        "unknown_evidence": [
            "Independent reconciliation does not establish publication/receipt timing, "
            "market completeness, "
            "source authority or lifecycle effective-versus-announcement semantics.",
            "Cito, full history v1 and optional supply classifications have no independent "
            "contract in this auditor; only explicitly bounded history v2 is covered.",
        ],
    }
    mismatches = report["mismatches"]

    def check(kind: str, key: str, expected: Any, actual: Any) -> bool:
        if expected == actual:
            return True
        mismatches.append({"kind": kind, "key": key, "expected": expected, "actual": actual})
        return False

    deliveries = {d["delivery_id"]: d for d in tables["source_delivery"]}
    original_hashes = {
        str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in original_paths or []
    }
    report["original_hashes"] = original_hashes
    basis_dependencies = _basis_dependency_ids(tables["import_batch"])
    primary_deliveries = {batch["delivery_id"] for batch in tables["import_batch"]}
    for delivery_id, delivery in deliveries.items():
        p = (archive / delivery["archive_location"]).resolve()
        if not p.is_relative_to(archive.resolve()) or not p.is_file():
            check("archive_unavailable", delivery_id, "existing archive member", None)
            continue
        check(
            "archive_hash",
            delivery_id,
            delivery["content_hash"],
            hashlib.sha256(p.read_bytes()).hexdigest(),
        )
        check("archive_size", delivery_id, delivery["byte_count"], p.stat().st_size)
        if original_paths is not None and (
            delivery_id not in basis_dependencies or delivery_id in primary_deliveries
        ):
            check(
                "original_hash_match",
                delivery_id,
                True,
                delivery["content_hash"] in original_hashes.values(),
            )
        check(
            "delivery_identity",
            delivery_id,
            generated_id(
                "delivery", {"source": delivery["source"], "receipt_key": delivery["receipt_key"]}
            ),
            delivery_id,
        )
        metadata = {
            key: delivery[key]
            for key in (
                "source",
                "receipt_key",
                "original_filename",
                "original_uri",
                "edition",
                "phase",
                "current_period",
                "target_period",
                "content_hash",
            )
        }
        for key in ("coverage", "replacement"):
            metadata[key] = json.loads(delivery[key])
        for prefix in ("received", "published"):
            metadata[prefix] = [
                delivery[prefix + "_" + suffix]
                for suffix in ("lower", "upper", "precision", "basis", "timezone")
            ]
        check("delivery_metadata_hash", delivery_id, digest(metadata), delivery["metadata_hash"])

    expected_packages = set()
    expected_mappings = {}
    for batch in tables["import_batch"]:
        iid = batch["import_id"]
        delivery = deliveries[batch["delivery_id"]]
        actual = [
            (table, row)
            for table in FIELDS
            for row in tables.get(table, [])
            if row["import_id"] == iid
        ]
        actual_issues = [r for r in tables["import_issue"] if r["import_id"] == iid]
        info = {
            "import_id": iid,
            "delivery_id": batch["delivery_id"],
            "parser_version": batch["parser_version"],
            "actual_assertions": len(actual),
            "actual_issues": len(actual_issues),
            "import_record_hash": persisted_record_hash(batch, ("processing_at", "completed_at")),
            "context_unknowns": {
                "publication_timing": delivery["published_precision"] == "unknown",
                "receipt_timing": delivery["received_precision"] == "unknown",
                "current_report_period": delivery["current_period"] is None,
                "market_completeness": "not_independently_established",
            },
        }
        report["imports"].append(info)
        start = len(mismatches)
        context = json.loads(batch["interpretation_context"])
        check("context_hash", iid, digest(context), batch["interpretation_context_hash"])
        identity_fields = {
            k: batch[k]
            for k in (
                "delivery_id",
                "parser_version",
                "schema_version",
                "mapping_version",
                "interpretation_context_hash",
            )
        }
        check("import_identity", iid, generated_id("import", identity_fields), iid)
        if batch["parser_version"] == "f3-identity-evidence-v1":
            try:
                document = json.loads(
                    (archive / delivery["archive_location"]).read_text(encoding="utf-8")
                )
                if (
                    document.get("version") != 1
                    or document.get("kind") != "f3_identity_evidence"
                    or context != {"version": 1, "kind": "identity_evidence"}
                    or batch["schema_version"] != "f3-identity-evidence-v1"
                ):
                    raise ValueError("independent_identity_document_schema_mismatch")
                entries = document["entries"]
                for support_id in document.get("supporting_delivery_ids", []):
                    check(
                        "identity_support_delivery",
                        iid + ":" + support_id,
                        True,
                        support_id in deliveries,
                    )
                for number, entry in enumerate(entries):
                    canonical_id = entry["canonical_identifier"]
                    if not re.fullmatch(r"[0-9]{6}", canonical_id):
                        raise ValueError("independent_identity_document_identifier")
                    pid = entry["package_id"]
                    derivation = entry.get("package_id_derivation")
                    if derivation:
                        check(
                            "generated_identity_package",
                            iid + ":" + canonical_id,
                            generated_id(derivation["kind"], derivation["value"]),
                            pid,
                        )
                    else:
                        info.setdefault("explicit_document_assertions", []).append(
                            "package_id:" + pid
                        )
                    expected_packages.add(pid)
                    available = entry["available"]
                    m = {
                        "package_id": pid,
                        "namespace": entry["namespace"],
                        "identifier_value": canonical_id,
                        "canonical_value": canonical_id,
                        "validity_kind": entry["validity_kind"],
                        "valid_from": entry.get("valid_from"),
                        "valid_to": entry.get("valid_to"),
                        "scope_key": entry.get("scope_key"),
                        "evidence_import_id": iid,
                        "source_locator": entry["source_locator"],
                        "mapping_version": entry["mapping_version"],
                    }
                    timing = [
                        available.get(k)
                        for k in ("lower", "upper", "precision", "basis", "timezone")
                    ]
                    mid = generated_id("mapping", {**m, "available": timing})
                    expected_mappings[(iid, str(number))] = {
                        "mapping_id": mid,
                        **m,
                        **{
                            "available_" + k: available.get(k)
                            for k in ("lower", "upper", "precision", "basis", "timezone")
                        },
                    }
                    supporting_rows = [
                        r
                        for table in FIELDS
                        for r in tables.get(table, [])
                        if r["canonical_identifier"] == canonical_id
                        and r["delivery_id"] in document.get("supporting_delivery_ids", [])
                    ]
                    check(
                        "identity_source_corroboration",
                        iid + ":" + canonical_id,
                        True,
                        bool(supporting_rows),
                    )
                    for support_oid in entry.get("supporting_observation_ids", []):
                        check(
                            "identity_support_observation",
                            iid + ":" + support_oid,
                            True,
                            any(r["observation_id"] == support_oid for r in supporting_rows),
                        )
                check("identity_document_observations", iid, [], actual)
                check("identity_document_issues", iid, [], actual_issues)
                for key, value in (
                    ("record_count", len(entries)),
                    ("accepted_count", len(entries)),
                    ("rejected_count", 0),
                ):
                    check("identity_document_count", iid + ":" + key, value, batch[key])
                info.update(
                    status="verified" if len(mismatches) == start else "quarantined",
                    expected_assertions=0,
                    expected_mapping_assertions=len(entries),
                    expected_issues=0,
                    mismatch_count=len(mismatches) - start,
                )
            except (ValueError, KeyError, OSError, TypeError) as exc:
                info.update(status="quarantined", reason=str(exc))
            continue
        for key in ("current_period", "target_period"):
            check(
                "context_delivery_binding",
                iid + ":" + key,
                delivery[key],
                context.get("adapter_context", {}).get(key),
            )
        try:
            expected_versions = {context["adapter"] + "-workbook-v1"}
            if context.get("adapter") == "lms-history":
                expected_versions = {"lms-history-workbook-v2"}
            elif context.get("adapter") == "medicine-prices":
                expected_versions = {"medicine-prices-pdf-v1"}
            elif context.get("adapter") == "lms-ftp":
                expected_versions = {"lms-ftp-zip-v1"}
            elif context.get("adapter") in ("introductions", "discontinued"):
                expected_versions.add(context["adapter"] + "-workbook-v2")
            if context.get("adapter") == "price-changes":
                expected_versions.add("price-changes-workbook-v2")
            elif context.get("adapter") == "introductions":
                expected_versions.add("introductions-workbook-v3")
            if context.get("adapter") in ("price-changes", "introductions", "discontinued"):
                expected_versions.add(context["adapter"] + "-rtf-v1")
            if (
                batch["parser_version"] not in expected_versions
                or batch["schema_version"] != "typed-observations-v1"
            ):
                raise ValueError("independent_version_contract_unavailable")
            if context["mapping_policy"] != "delivery" or batch["validation_status"] != "accepted":
                raise ValueError("independent_publication_contract_unavailable")
            basis_profile = (
                None
                if batch["parser_version"] == "lms-ftp-zip-v1"
                and "price_basis_override" not in context["adapter_context"]
                else _basis_override(
                    context["adapter_context"],
                    enabled=batch["parser_version"]
                    in {"price-changes-workbook-v2", "introductions-workbook-v3", "lms-ftp-zip-v1"},
                )
            )
            if basis_profile is not None:
                info["basis_context_evidence"] = _basis_evidence_pin(
                    basis_profile, deliveries, archive
                )
            source_bytes = (archive / delivery["archive_location"]).read_bytes()
            if batch["parser_version"] == "lms-ftp-zip-v1":
                from .lms_ftp_audit import expected_lms_ftp, validate_delivery

                parsed = expected_lms_ftp(
                    source_bytes, context["adapter"], context["adapter_context"]
                )
                validate_delivery(
                    delivery,
                    context["adapter_context"],
                    tables.get("price_period", []),
                    parsed["source_rows"],
                )
            elif batch["parser_version"] == "medicine-prices-pdf-v1":
                from .pdf_audit import expected_pdf

                parsed = expected_pdf(source_bytes, context["adapter"], context["adapter_context"])
            elif batch["parser_version"].endswith("-rtf-v1"):
                from .rtf_audit import expected_rtf

                parsed = expected_rtf(source_bytes, context["adapter"], context["adapter_context"])
            else:
                parsed = expected_report(
                    source_bytes,
                    context["adapter"],
                    context["adapter_context"],
                    parser_version=batch["parser_version"],
                )
        except (ValueError, OSError, KeyError) as exc:
            info.update(status="quarantined", reason=str(exc))
            report["quarantined_observation_ids"].extend(r["observation_id"] for _, r in actual)
            continue
        info.update(
            source_rows=parsed["source_rows"],
            independent_schema={
                k: parsed[k] for k in ("sheet", "header_row", "metadata", "date_system")
            },
        )
        if "physical_accounting" in parsed:
            info["physical_accounting"] = parsed["physical_accounting"]
        if "scope_accounting" in parsed:
            info["scope_accounting"] = parsed["scope_accounting"]
            coverage = json.loads(delivery["coverage"])
            check("subset_coverage_kind", iid, "partial_extract", coverage.get("kind"))
            expected_filter = {
                key: context["adapter_context"][key]
                for key in ("subset_identifiers", "subset_periods")
            }
            check("subset_coverage_filter", iid, expected_filter, coverage.get("filter"))
        candidates = parsed["assertions"]
        # Independent grouping retains every physical assertion and compares field states,
        # rather than coalescing duplicate rows or using last-row order.
        groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
        business_keys = {
            "price_observation": (
                "price_type",
                "currency",
                "price_basis",
                "tax_basis",
                "fee_basis",
                "price_role",
            ),
            "package_observation": (),
            "supply_observation": (
                "observation_date",
                "wholesaler_namespace",
                "wholesaler_identifier",
            ),
            "lifecycle_event": ("event_type",),
            "classification_observation": ("classification_context",),
        }
        for a in candidates:
            business_key = [
                a[k]
                for k in (
                    "table",
                    "canonical_identifier",
                    "effective_period",
                    "report_period",
                    "effective_date",
                )
            ]
            groups[
                canonical(business_key + [a["values"].get(k) for k in business_keys[a["table"]]])
            ].append(a)
        duplicate_groups = 0
        for group in groups.values():
            if len(group) < 2:
                continue
            duplicate_groups += 1
            for field in FIELDS[group[0]["table"]]:
                values = set()
                for a in group:
                    state = a["presence"].get(
                        field, "present" if a["values"].get(field) is not None else "unreported"
                    )
                    if state in ("present", "blank", "withdrawn", "conflicting"):
                        values.add(canonical([state, a["values"].get(field)]))
                if len(values) > 1:
                    for a in group:
                        state = a["presence"].get(
                            field, "present" if a["values"].get(field) is not None else "unreported"
                        )
                        if state in ("present", "blank", "withdrawn", "conflicting"):
                            a["presence"][field] = "conflicting"
                            a["quality"] = "conflicting"
                            parsed["issues"].append(
                                {
                                    "source_locator": a["source_locator"],
                                    "code": "conflicting_assertions",
                                    "identifier": a["canonical_identifier"],
                                    "field": field,
                                    "period": a["effective_period"],
                                    "severity": "error",
                                    "blocking_scope": "field",
                                }
                            )
        expected_records = []
        for a in candidates:
            pid = generated_id(
                "source-package",
                {
                    "delivery_id": batch["delivery_id"],
                    "namespace": "dk:varenummer",
                    "identifier": a["canonical_identifier"],
                },
            )
            expected_packages.add(pid)
            map_key = batch["delivery_id"], a["canonical_identifier"]
            if map_key not in expected_mappings:
                m = {
                    "package_id": pid,
                    "namespace": "dk:varenummer",
                    "identifier_value": a["canonical_identifier"],
                    "canonical_value": a["canonical_identifier"],
                    "validity_kind": "scope",
                    "valid_from": None,
                    "valid_to": None,
                    "scope_key": "delivery:" + batch["delivery_id"],
                    "evidence_import_id": iid,
                    "source_locator": a["source_locator"],
                    "mapping_version": batch["mapping_version"],
                }
                mid = generated_id(
                    "mapping", {**m, "available": [None, None, "unknown", None, None]}
                )
                expected_mappings[map_key] = {
                    "mapping_id": mid,
                    **m,
                    "available_lower": None,
                    "available_upper": None,
                    "available_precision": "unknown",
                    "available_basis": None,
                    "available_timezone": None,
                }
            presence = {
                k: a["presence"].get(
                    k, "present" if a["values"].get(k) is not None else "unreported"
                )
                for k in FIELDS[a["table"]]
            }
            presence.update(a["presence"])
            quality = a["quality"]
            if quality == "valid":
                if "invalid" in presence.values():
                    quality = "invalid"
                elif set(presence.values()) & {"blank", "missing", "withdrawn"}:
                    quality = "incomplete"
            record = {
                "import_id": iid,
                "delivery_id": batch["delivery_id"],
                "package_id": pid,
                "mapping_id": expected_mappings[map_key]["mapping_id"],
                "identifier_namespace": "dk:varenummer",
                **{
                    k: a[k]
                    for k in (
                        "source_locator",
                        "source_identifier",
                        "canonical_identifier",
                        "effective_period",
                        "report_period",
                        "effective_date",
                    )
                },
                "quality_state": quality,
                "field_presence": canonical(presence),
                "source_values": canonical(a["raw"]),
                **{k: a["values"].get(k) for k in FIELDS[a["table"]]},
            }
            oid = generated_id("observation", {"table": a["table"], **record})
            expected_records.append((a["table"], {"observation_id": oid, **record}))
        actual_by_loc = {(t, r["source_locator"]): r for t, r in actual}
        if len(actual_by_loc) != len(actual):
            check("duplicate_stored_locator", iid, len(actual), len(actual_by_loc))
        for table, expected in expected_records:
            location_key = table, expected["source_locator"]
            stored = actual_by_loc.pop(location_key, None)
            if check("observation_all_fields", iid + ":" + str(location_key), expected, stored):
                report["verified_observation_ids"].append(expected["observation_id"])
            elif stored:
                report["quarantined_observation_ids"].append(stored["observation_id"])
        for location_key, stored in actual_by_loc.items():
            check("unexpected_observation", iid + ":" + str(location_key), None, stored)
            report["quarantined_observation_ids"].append(stored["observation_id"])
        expected_issues = {}
        for issue in parsed["issues"]:
            record = {"import_id": iid, **issue}
            issue_id = generated_id("issue", record)
            expected_issues[issue_id] = {"issue_id": issue_id, **record}
        check(
            "issues_all_fields",
            iid,
            sorted(expected_issues.values(), key=canonical),
            sorted(actual_issues, key=canonical),
        )
        for key, value in (
            ("record_count", len(candidates)),
            ("accepted_count", len(candidates)),
            ("rejected_count", 0),
        ):
            check("assertion_count", iid + ":" + key, value, batch[key])
        info.update(
            expected_assertions=len(candidates),
            expected_issues=len(expected_issues),
            duplicate_business_key_groups=duplicate_groups,
            mismatch_count=len(mismatches) - start,
            status="verified" if len(mismatches) == start else "quarantined",
        )
    check(
        "package_all_records",
        "all",
        sorted(expected_packages),
        sorted(r["package_id"] for r in tables["package"]),
    )
    check(
        "mapping_all_fields",
        "all",
        sorted(expected_mappings.values(), key=canonical),
        sorted(tables["package_identifier"], key=canonical),
    )
    verified_ids = set(report["verified_observation_ids"])
    report["observation_hashes"] = {
        r["observation_id"]: digest(r)
        for table in FIELDS
        for r in tables.get(table, [])
        if r["observation_id"] in verified_ids
    }
    report["verified_count"] = len(report["verified_observation_ids"])
    report["quarantined_count"] = len(report["quarantined_observation_ids"])
    report["complete"] = (
        not mismatches
        and not report["quarantined_observation_ids"]
        and all(i["status"] == "verified" for i in report["imports"])
    )
    report["mapping_hashes"] = (
        {
            row["mapping_id"]: persisted_record_hash(row, ("available_lower", "available_upper"))
            for row in tables["package_identifier"]
        }
        if report["complete"]
        else {}
    )
    return report


class _PrivateArgumentParser(argparse.ArgumentParser):
    def error(self, message: str) -> Never:
        # Argparse's default diagnostic can echo arbitrary private argument text.
        self.exit(2, "source-audit: invalid arguments\n")


def main(argv: list[str] | None = None) -> int:
    parser = _PrivateArgumentParser(
        prog="source-audit", description="Private read-only original reconciliation"
    )
    parser.add_argument("--database", type=Path, required=True)
    parser.add_argument("--archive", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--original-manifest", type=Path)
    parser.add_argument("--import-id", action="append", help="Audit every row of explicit imports")
    args = parser.parse_args(argv)
    try:
        output = external_path(args.output)
        handle = output.open("x", encoding="utf-8")
    except Exception:
        print(canonical({"complete": False, "error": "audit_output_unavailable"}), file=sys.stderr)
        return 2
    try:
        with handle:
            try:
                database = external_path(args.database)
                archive = external_path(args.archive)
                originals = None
                if args.original_manifest:
                    manifest = external_path(args.original_manifest)
                    originals = [
                        external_path(Path(p))
                        for p in json.loads(manifest.read_text(encoding="utf-8"))["source_hashes"]
                    ]
                report = reconcile(database, archive, originals, import_ids=args.import_id)
            except Exception as exc:
                json.dump(
                    {
                        "version": 1,
                        "complete": False,
                        "audit_error": {"type": type(exc).__name__, "detail": str(exc)},
                    },
                    handle,
                    ensure_ascii=True,
                    indent=2,
                )
                print(canonical({"complete": False, "error": "audit_failed"}), file=sys.stderr)
                return 2
            json.dump(report, handle, ensure_ascii=True, indent=2)
    except Exception:
        # Even failure to persist the private diagnostic must not expose its text.
        print(canonical({"complete": False, "error": "audit_output_failed"}), file=sys.stderr)
        return 2
    print(canonical({k: report[k] for k in ("complete", "verified_count", "quarantined_count")}))
    return 0 if report["complete"] is True else 1


if __name__ == "__main__":
    raise SystemExit(main())
