"""Independent synthetic examples for the concise decision workbook."""

import json
from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from bef_redbird.domain import (
    AnalysisContext,
    AnalysisSettings,
    CitoItem,
    CitoSnapshot,
    HistoryDataset,
    QueueSettings,
    Review,
    ReviewState,
    SourceInfo,
)
from bef_redbird.export import export_analysis, export_ids
from bef_redbird.service import build_analysis


@pytest.fixture
def decision_analysis():
    periods = (date(2025, 1, 6), date(2025, 1, 20), date(2025, 2, 3), date(2025, 2, 17))
    snapshot = CitoSnapshot(
        SourceInfo("synthetic-cito.xlsx", "a" * 64, 10),
        {
            "000001": CitoItem("000001", "", 10001, "A", "C", "00007", (2,)),
            "000002": CitoItem("000002", "", 9999, "B", "C", "00007", (3,)),
            "000003": CitoItem("000003", "H", 2500, "A", "A", "00007", (4,)),
            "000004": CitoItem("000004", "", 5000, "A", "C", None, (5,)),
        },
        [],
        4,
        4,
        0,
        incomplete_groups={"00007"},
    )
    history = HistoryDataset(
        SourceInfo("synthetic-lms.xlsx", "b" * 64, 20),
        (*periods, date(2025, 3, 3)),
        {
            # Median of four observable values is 15001.5 øre. Future price is excluded.
            "000001": (10001, 20002, 20002, 10001, 999999),
            "000002": (10001, 20002, 20002, 10001, 999999),
        },
        {
            "000001": "Syntetisk pakning, 5 mg, tabletter, 100 stk., Eksempelfirma",
            "000002": "Anden syntetisk pakning",
        },
        {"000001": (2,), "000002": (3,)},
        [],
    )
    return build_analysis(snapshot, history, AnalysisContext(periods[-1]), AnalysisSettings(4, 4))


def _book(analysis, state, scope="queue"):
    return load_workbook(BytesIO(export_analysis(analysis, state, scope)), data_only=False)


def test_concise_overview_has_independent_exact_prices_and_prominent_review(decision_analysis):
    state = ReviewState(
        decision_analysis.analysis_id,
        reviews={"000001": Review("Undersøg nærmere", "Afklar efterspørgsel")},
    )
    book = _book(decision_analysis, state)
    assert book.sheetnames == [
        "Indkøbsoverblik",
        "Kandidater",
        "Gruppeoverblik",
        "Metode og datakontrol",
    ]
    overview = book.worksheets[0]
    assert overview.max_column == 7
    row = list(overview.values)[3]
    assert row[:6] == (
        "000001",
        "Syntetisk pakning, 5 mg, tabletter, 100 stk., Eksempelfirma",
        100.01,
        100.01,
        150.015,
        "Undersøg nærmere",
    )
    assert row[6] == "Afklar efterspørgsel\n\nA til C.\nUafklarede kilderækker i gruppen."
    assert book["Kandidater"]["N2"].value == 50.005
    assert "historisk" in overview.cell(3, 5).value.lower()
    assert "17-02-2025" in overview["A1"].value
    assert "06-01-2025 til 17-02-2025, 4 perioder" in overview["C1"].value
    assert "ikke en besparelse" in overview["A2"].value
    for col in (3, 4, 5):
        assert overview.cell(4, col).data_type == "n"
    assert overview["A4"].data_type == "s"


def test_clean_context_only_shows_supplied_transition(decision_analysis):
    decision_analysis.snapshot.incomplete_groups.clear()
    book = _book(decision_analysis, ReviewState(decision_analysis.analysis_id))
    assert book["Indkøbsoverblik"]["G4"].value == "A til C."


@pytest.mark.parametrize("scope", ["queue", "selected", "all_eligible"])
def test_scope_order_parity_detail_and_manifest_unchanged(decision_analysis, scope):
    state = ReviewState(
        decision_analysis.analysis_id,
        QueueSettings(top_n=1),
        {"000002": Review("Fravalgt", "Kontrollér prisdefinition")},
        {"000002", "000003"},
    )
    expected_ids = export_ids(decision_analysis, state, scope)
    book = _book(decision_analysis, state, scope)
    overview = list(book["Indkøbsoverblik"].values)[3:]
    details = list(book["Kandidater"].values)
    assert tuple(row[0] for row in overview) == expected_ids
    assert tuple(row[1] for row in details[1:]) == expected_ids
    assert len(details[0]) == 29
    assert details[0][2:8] == (
        "Opbevaringskode",
        "Kostpris (DKK)",
        "Nuværende kode",
        "Kommende kode",
        "Gruppe",
        "Aktuel AIP (DKK)",
    )
    for short, full in zip(overview, details[1:], strict=True):
        assert short[0] == full[1]
        assert short[5] == full[25]
        if full[26]:
            assert short[6].startswith(full[26] + "\n\n")
        assert short[2:5] == (full[3], full[7], full[12])
    metadata = list(book["Metode og datakontrol"].values)[1:]
    manifest = dict((key, value) for key, value in metadata)
    for key, value in decision_analysis.manifest().items():
        assert manifest[key] == (
            json.dumps(value, ensure_ascii=False, sort_keys=True)
            if isinstance(value, dict)
            else value
        )
    assert [v for k, v in metadata if k == "ordered_exported_id"] == list(expected_ids)
    assert len([v for k, v in metadata if k == "selected_group_relation"]) == len(expected_ids)
    assert json.loads(manifest["queue_export"]) == {
        "mode": "top_n",
        "top_n": 1,
        "effective_n": 1,
        "scope": scope,
        "exported_count": len(expected_ids),
    }


def test_unranked_mismatch_missing_and_unfiltered_alternatives_remain_visible(decision_analysis):
    book = _book(decision_analysis, ReviewState(decision_analysis.analysis_id), "all_eligible")
    overview = {row[0]: row for row in list(book["Indkøbsoverblik"].values)[3:]}
    assert overview["000002"][2:5] == (99.99, 100.01, 150.015)
    assert "stemmer ikke overens" in overview["000002"][6]
    assert "B til C" in overview["000002"][6]
    assert overview["000004"][1] == "Varenummer 000004"
    assert overview["000004"][3:5] == (None, None)
    assert "Gruppe mangler" in overview["000004"][6]
    assert "AIP mangler" in overview["000004"][6]
    groups = {row[0]: row for row in list(book["Gruppeoverblik"].values)[1:]}
    assert set(groups) == {"000001", "000002", "000003", "000004"}
    assert groups["000003"][8:10] == ("Ja", "Mangler")
    assert groups["000003"][2:4] == ("Varenummer 000003", "Nej")
    assert groups["000001"][2:4] == (decision_analysis.history.labels["000001"], "Ja")


@pytest.mark.parametrize(
    "text", ["=1+1", "+SUM(A1)", "-1+2", "@A1", "\t=2", "\n=3", "https://example.invalid"]
)
def test_new_labels_notes_and_context_are_plain_text(decision_analysis, text):
    decision_analysis.history.labels["000001"] = text
    state = ReviewState(decision_analysis.analysis_id, reviews={"000001": Review(note=text)})
    book = _book(decision_analysis, state)
    for sheet, coord in (
        ("Indkøbsoverblik", "B4"),
        ("Kandidater", "AA2"),
        ("Gruppeoverblik", "C2"),
    ):
        cell = book[sheet][coord]
        assert cell.value == text
        assert cell.data_type == "s"
        assert cell.hyperlink is None
    assert book["Indkøbsoverblik"]["G4"].value.startswith(text + "\n\n")
    assert book["Indkøbsoverblik"]["G4"].data_type == "s"
    assert not any(
        cell.data_type == "f" or cell.hyperlink is not None
        for sheet in book
        for row in sheet
        for cell in row
    )


def test_readable_structure_freezes_identifiers_and_retains_print_headers(decision_analysis):
    book = _book(decision_analysis, ReviewState(decision_analysis.analysis_id))
    overview = book["Indkøbsoverblik"]
    assert overview.freeze_panes == "C4"
    assert overview.auto_filter.ref == "A3:G4"
    assert overview.column_dimensions["B"].width >= 26
    assert overview.column_dimensions["G"].width >= 25
    assert overview["B4"].alignment.wrap_text
    assert overview["G4"].alignment.wrap_text
    assert overview["F4"].fill.fgColor != overview["B4"].fill.fgColor
    for sheet in book:
        assert not sheet.merged_cells.ranges
        assert sheet.sheet_view.showGridLines is False
        assert sheet.page_setup.orientation == "landscape"
        assert sheet.print_title_rows
        assert sheet.print_title_cols == "$A:$B"
        assert sheet.auto_filter.ref
        assert sheet.print_area
        assert not sheet["A1"].alignment.shrink_to_fit
        assert sheet.row_dimensions[3 if sheet == overview else 1].height >= 40
    assert book["Kandidater"]["K2"].number_format == "0.0%"
    assert book["Metode og datakontrol"].freeze_panes == "B2"
