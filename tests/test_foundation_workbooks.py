"""Independently expected synthetic acceptance for workbook assertion adapters."""

from __future__ import annotations

import io
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

from bef_redbird.foundation.adapters.base import AdapterError
from bef_redbird.foundation.adapters.workbooks import parse_cito, parse_history

CURRENT = date(2025, 1, 20)
TARGET = date(2025, 2, 3)
CITO_HEADERS = ["varenr", "opbevkode", "kostpris", "subkode", "nysubkode", "subgrp"]
CITO_CONTEXT: dict[str, Any] = {
    "current_period": CURRENT.isoformat(),
    "target_period": TARGET.isoformat(),
    "price_type": "aip",
    "currency": "DKK",
    "price_basis": "per_package",
    "tax_basis": "unknown",
    "fee_basis": "unknown",
    "classification_context": "ordinary",
    "group_namespace": "synthetic-cito",
}
HISTORY_CONTEXT: dict[str, Any] = {
    "current_period": CURRENT.isoformat(),
    "tax_basis": "unknown",
    "fee_basis": "unknown",
}


def workbook(tmp_path: Path, sheet: str, rows: list[list[Any]], *, biff: bool = False) -> bytes:
    """Write real files to exercise bytes read back from both supported containers."""
    if biff:
        import xlwt

        book = xlwt.Workbook()
        legacy_sheet = book.add_sheet(sheet)
        for rowno, row in enumerate(rows):
            for col, value in enumerate(row):
                if value is not None:
                    legacy_sheet.write(rowno, col, value)
        path = tmp_path / "synthetic.xls"
        book.save(str(path))
    else:
        book = Workbook()
        target = book.active
        assert target is not None
        target.title = sheet
        for row in rows:
            target.append(row)
        path = tmp_path / "synthetic.xlsx"
        book.save(path)
        book.close()
    return path.read_bytes()


@pytest.mark.parametrize("biff", [False, True])
def test_cito_both_formats_preserve_assertions_identity_and_group_scope(
    tmp_path: Path, biff: bool
) -> None:
    data = workbook(
        tmp_path,
        "Invented legacy" if biff else "Table",
        [
            [" Ny Subkode ", "SubGrp", "KOSTPRIS", "VARENR", "SUBKODE", "OPBEVKODE", "private"],
            ["C", "0007", "1.234,56", "000123", "a", "h", "never retained"],
            ["C", "0007", "1.234,56", "000123", "a", "h", "never retained"],
            ["B", 9, 0, 24, "B", None, "never retained"],
        ],
        biff=biff,
    )
    parsed = parse_cito(data, CITO_CONTEXT)
    assert parsed.records == 3
    assert not parsed.issues
    assert len(parsed.assertions) == 12
    prices = [a for a in parsed.assertions if a.table == "price_observation"]
    assert [a.values["amount_ore"] for a in prices] == [123456, 123456, 0]
    assert prices[0].source_identifier == "000123"
    assert prices[2].source_identifier == ("24.0" if biff else "24")
    assert prices[2].raw["source_identifier"] == ("24.0" if biff else "24")
    assert prices[0].effective_period == CURRENT
    assert prices[0].report_period == CURRENT
    sheet = "Invented legacy" if biff else "Table"
    assert prices[0].source_locator == f"sheet:{sheet}!C2"
    classifications = [a for a in parsed.assertions if a.table == "classification_observation"]
    assert classifications[0].values["abc_code"] == "A"
    assert classifications[0].values["group_code"] == "0007"
    assert classifications[0].values["group_scope"] == "2025-01-20"
    assert classifications[1].effective_period == TARGET
    assert classifications[1].report_period == CURRENT
    assert classifications[1].values["group_code"] is None
    assert classifications[1].values["group_scope"] is None
    assert classifications[1].presence["group_code"] == "unreported"
    packages = [a for a in parsed.assertions if a.table == "package_observation"]
    assert packages[0].source_locator == f"sheet:{sheet}!row:2"
    assert packages[0].values["storage_code"] == "H"
    assert packages[0].presence["name"] == "unreported"
    assert packages[2].presence["storage_code"] == "blank"
    assert "private" not in repr(parsed)
    assert "never retained" not in repr(parsed)


@pytest.mark.parametrize(
    "value,code",
    [
        ("1,001", "unsupported_money_precision"),
        (12.3400000001, "unsupported_money_precision"),
        (-1, "negative_money"),
        ("12.34", "malformed_money"),
        ("truncated,", "malformed_money"),
    ],
)
def test_cito_exact_money_is_scoped_and_raw_retained(tmp_path: Path, value: Any, code: str) -> None:
    data = workbook(tmp_path, "Table", [CITO_HEADERS, ["000123", "", value, "A", "C", "7"]])
    parsed = parse_cito(data, CITO_CONTEXT)
    assert len(parsed.assertions) == 4
    price = parsed.assertions[0]
    assert price.values["amount_ore"] is None
    assert price.presence == {"amount_ore": "invalid"}
    assert price.quality_state == "invalid"
    assert price.raw["kostpris"] == str(value)
    assert [(i.code, i.source_locator, i.identifier, i.field, i.period) for i in parsed.issues] == [
        (code, "sheet:Table!C2", "000123", "amount_ore", CURRENT),
    ]
    assert parsed.assertions[2].values["abc_code"] == "A"


@pytest.mark.parametrize("biff", [False, True])
def test_cito_formula_is_invalid_only_in_affected_field(tmp_path: Path, biff: bool) -> None:
    value: Any = "=1+1"
    if biff:
        import xlwt

        value = xlwt.Formula("1+1")
    data = workbook(tmp_path, "Table", [CITO_HEADERS, ["123", "", value, "A", "C", 7]], biff=biff)
    parsed = parse_cito(data, CITO_CONTEXT)
    assert parsed.assertions[0].presence["amount_ore"] == "invalid"
    assert [i.code for i in parsed.issues] == ["formula_cell"]
    assert parsed.assertions[2].presence["abc_code"] == "present"


@pytest.mark.parametrize(
    "key",
    [
        "current_period",
        "target_period",
        "price_type",
        "currency",
        "price_basis",
        "tax_basis",
        "fee_basis",
        "classification_context",
        "group_namespace",
    ],
)
def test_cito_requires_explicit_context(tmp_path: Path, key: str) -> None:
    data = workbook(tmp_path, "Table", [CITO_HEADERS, [123, "", 1, "A", "C", 7]])
    context = {k: v for k, v in CITO_CONTEXT.items() if k != key}
    with pytest.raises(AdapterError):
        parse_cito(data, context)


def test_cito_current_and_target_must_differ(tmp_path: Path) -> None:
    data = workbook(tmp_path, "Table", [CITO_HEADERS])
    with pytest.raises(AdapterError, match="identical_cito_periods"):
        parse_cito(data, {**CITO_CONTEXT, "target_period": CURRENT.isoformat()})


def test_cito_upcoming_period_cannot_precede_current(tmp_path: Path) -> None:
    data = workbook(tmp_path, "Table", [CITO_HEADERS])
    with pytest.raises(AdapterError, match="reversed_cito_periods"):
        parse_cito(data, {**CITO_CONTEXT, "target_period": "2025-01-06"})


@pytest.mark.parametrize("value", [CURRENT, datetime(2025, 1, 20), 20250120, True])
@pytest.mark.parametrize("key", ["current_period", "target_period"])
def test_cito_context_dates_require_explicit_iso_text(tmp_path: Path, value: Any, key: str) -> None:
    data = workbook(tmp_path, "Table", [CITO_HEADERS])
    with pytest.raises(AdapterError, match=f"invalid_{key}_context"):
        parse_cito(data, {**CITO_CONTEXT, key: value})


@pytest.mark.parametrize("biff", [False, True])
@pytest.mark.parametrize("value", [True, False])
@pytest.mark.parametrize(
    "col,code", [(0, "invalid_identifier"), (2, "malformed_money"), (5, "invalid_group_code")]
)
def test_boolean_cells_cannot_become_numeric_business_values(
    tmp_path: Path,
    biff: bool,
    value: bool,
    col: int,
    code: str,
) -> None:
    row: list[Any] = ["000123", "", 1, "A", "C", "7"]
    row[col] = value
    data = workbook(tmp_path, "Table", [CITO_HEADERS, row], biff=biff)
    parsed = parse_cito(data, CITO_CONTEXT)
    assert [issue.code for issue in parsed.issues] == [code]
    if col == 0:
        assert parsed.assertions == []
    elif col == 2:
        price = parsed.assertions[0]
        assert price.values["amount_ore"] is None
        assert price.presence["amount_ore"] == "invalid"
        assert price.raw["kostpris"] is value
        assert price.raw["kostpris_cell_type"] == "b"
    else:
        current_classification = parsed.assertions[2]
        assert current_classification.values["group_code"] is None
        assert current_classification.presence["group_code"] == "invalid"
        assert current_classification.raw["subgrp"] is value
        assert current_classification.raw["subgrp_cell_type"] == "b"


@pytest.mark.parametrize("biff", [False, True])
@pytest.mark.parametrize("col", [0, 2, 5])
def test_date_cells_cannot_become_numeric_business_values(
    tmp_path: Path, biff: bool, col: int
) -> None:
    row: list[Any] = ["000123", "", 1, "A", "C", "7"]
    if biff:
        import xlwt

        book = xlwt.Workbook()
        sheet = book.add_sheet("Table")
        for index, heading in enumerate(CITO_HEADERS):
            sheet.write(0, index, heading)
        for index, value in enumerate(row):
            if index != col:
                sheet.write(1, index, value)
        # Real BIFF date cell: its serial 1 could otherwise become package
        # 000001, 100 ore or group 1. Keep the original serial as evidence.
        sheet.write(1, col, 1, xlwt.easyxf(num_format_str="DD/MM/YYYY"))
        path = tmp_path / "date.xls"
        book.save(str(path))
        data = path.read_bytes()
    else:
        row[col] = datetime(1900, 1, 1)
        data = workbook(tmp_path, "Table", [CITO_HEADERS, row])
    parsed = parse_cito(data, CITO_CONTEXT)
    if col == 0:
        assert [issue.code for issue in parsed.issues] == ["invalid_identifier_cell"]
        assert parsed.assertions == []
    else:
        assert [issue.code for issue in parsed.issues] == ["invalid_date_cell"]
        assertion = parsed.assertions[0 if col == 2 else 2]
        field = "amount_ore" if col == 2 else "group_code"
        raw_field = "kostpris" if col == 2 else "subgrp"
        assert assertion.values[field] is None
        assert assertion.presence[field] == "invalid"
        assert assertion.raw[f"{raw_field}_cell_type"] == "d"
        assert assertion.raw[raw_field] == ("1.0" if biff else "1900-01-01T00:00:00")


def test_biff_formula_heading_is_rejected_instead_of_using_cached_header(tmp_path: Path) -> None:
    import xlwt

    data = workbook(
        tmp_path,
        "Table",
        [
            [*CITO_HEADERS, xlwt.Formula("1+1")],
            ["000123", "", 1, "A", "C", "7"],
        ],
        biff=True,
    )
    with pytest.raises(AdapterError, match="invalid_header_cell"):
        parse_cito(data, CITO_CONTEXT)


def test_cito_invalid_group_is_scoped_and_uses_safe_code(tmp_path: Path) -> None:
    data = workbook(
        tmp_path, "Table", [CITO_HEADERS, [123, "", 1, "A", "C", "invalid private group"]]
    )
    parsed = parse_cito(data, CITO_CONTEXT)
    assert [(i.code, i.field, i.source_locator) for i in parsed.issues] == [
        ("invalid_group_code", "group_code", "sheet:Table!F2"),
    ]
    assert parsed.assertions[2].presence["group_code"] == "invalid"
    assert parsed.assertions[3].presence["group_code"] == "unreported"


@pytest.mark.parametrize(
    "key,value", [("currency", "EUR"), ("price_type", "consumer"), ("price_basis", "per_unit")]
)
def test_history_rejects_context_that_conflicts_with_aip_source(
    tmp_path: Path, key: str, value: str
) -> None:
    data = workbook(
        tmp_path, "Medicinpriser data", [["Varenummer", "Indikator", "20250120"], [123, "AIP", 1]]
    )
    with pytest.raises(AdapterError, match="conflicting_history_price_context"):
        parse_history(data, {**HISTORY_CONTEXT, key: value})


def test_history_dates_duplicate_rows_and_current_attributes(tmp_path: Path) -> None:
    data = workbook(
        tmp_path,
        "Medicinpriser data",
        [
            [
                "Varenummer",
                "Indikator",
                "Lægemiddel",
                "Firma",
                datetime(2025, 1, 6),
                "20250120",
                "unrecognized outside",
            ],
            ["000123", "AIP", "invented name", "company one", "1,23", 2, "not retained"],
            ["000123", "AIP", None, "company two", "1,24", 2, "not retained"],
            ["000123", "AIP ", "not exact", "not exact", 888, 999],
            ["000123", "AUP", "not AIP", "not AIP", 888, 999],
        ],
    )
    parsed = parse_history(data, HISTORY_CONTEXT)
    assert parsed.records == 2
    assert not parsed.issues
    prices = [a for a in parsed.assertions if a.table == "price_observation"]
    assert [a.values["amount_ore"] for a in prices] == [123, 200, 124, 200]
    assert [a.effective_period for a in prices] == [date(2025, 1, 6), CURRENT] * 2
    assert prices[0].source_locator == "sheet:Medicinpriser data!E2"
    assert prices[0].values["price_type"] == "aip"
    assert prices[0].values["price_basis"] == "per_package"
    packages = [a for a in parsed.assertions if a.table == "package_observation"]
    assert [a.values["company"] for a in packages] == ["company one", "company two"]
    assert all(a.effective_period == CURRENT and a.report_period == CURRENT for a in packages)
    assert packages[0].presence["form"] == "unreported"
    assert packages[1].presence["name"] == "blank"
    assert "not retained" not in repr(parsed)


def test_history_unknown_attribute_period_is_not_backdated(tmp_path: Path) -> None:
    data = workbook(
        tmp_path,
        "Medicinpriser data",
        [
            ["Varenummer", "Indikator", "Firma", "20200106"],
            ["000001", "AIP", "present firm", 1],
        ],
    )
    parsed = parse_history(data, {"tax_basis": "unknown", "fee_basis": "unknown"})
    price, attributes = parsed.assertions
    assert price.effective_period == date(2020, 1, 6)
    assert attributes.effective_period is None
    assert attributes.report_period is None
    assert attributes.values["company"] == "present firm"
    assert [i.code for i in parsed.issues] == ["unknown_attribute_period"]


def test_history_invalid_and_blank_cells_remain_distinct(tmp_path: Path) -> None:
    data = workbook(
        tmp_path,
        "Medicinpriser data",
        [
            ["Varenummer", "Indikator", "Firma", "20250106", "20250120"],
            ["000001", "AIP", "=1+1", "1,001", None],
            ["000001", "AIP", "valid", 1, 0],
        ],
    )
    parsed = parse_history(data, HISTORY_CONTEXT)
    first, blank, attributes, second, zero, _ = parsed.assertions
    assert first.presence["amount_ore"] == "invalid"
    assert blank.presence["amount_ore"] == "blank"
    assert attributes.presence["company"] == "invalid"
    assert second.values["amount_ore"] == 100
    assert zero.values["amount_ore"] == 0
    assert [(i.code, i.field) for i in parsed.issues] == [
        ("unsupported_money_precision", "amount_ore"),
        ("formula_cell", "company"),
    ]


@pytest.mark.parametrize(
    "headings",
    [
        ["20250106", "20250106"],
        ["20250106", "unknown inside", "20250120"],
        ["20250106", None, "20250120"],
        ["2025-01-06"],
        ["20250230"],
        ["no dates"],
    ],
)
def test_history_malformed_date_structure_is_fatal(tmp_path: Path, headings: list[Any]) -> None:
    data = workbook(
        tmp_path,
        "Medicinpriser data",
        [["Varenummer", "Indikator", *headings], ["000001", "AIP", 1, 2, 3]],
    )
    with pytest.raises(AdapterError):
        parse_history(data, HISTORY_CONTEXT)


def test_history_reference_column_between_dates_is_supported(tmp_path: Path) -> None:
    data = workbook(
        tmp_path,
        "Medicinpriser data",
        [
            ["Varenummer", "Indikator", "20250106", "Firma", "20250120"],
            ["000001", "AIP", 1, "invented firm", 2],
        ],
    )
    parsed = parse_history(data, HISTORY_CONTEXT)
    assert [a.values["amount_ore"] for a in parsed.assertions[:2]] == [100, 200]
    assert parsed.assertions[-1].values["company"] == "invented firm"


def test_history_error_cell_and_identifier_cell_do_not_leak_cached_values(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Medicinpriser data"
    sheet.append(["Varenummer", "Indikator", "20250120"])
    sheet.append(["000001", "AIP", "#DIV/0!"])
    sheet.append(["=1+1", "AIP", 2])
    sheet["C2"].data_type = "e"
    output = io.BytesIO()
    book.save(output)
    parsed = parse_history(output.getvalue(), HISTORY_CONTEXT)
    assert len(parsed.assertions) == 1
    assert parsed.assertions[0].presence["amount_ore"] == "invalid"
    assert [i.code for i in parsed.issues] == ["error_cell", "invalid_identifier_cell"]


@pytest.mark.parametrize(
    "parser,context", [(parse_cito, CITO_CONTEXT), (parse_history, HISTORY_CONTEXT)]
)
def test_invalid_container_is_fatal_without_source_content(
    parser: Any, context: dict[str, Any]
) -> None:
    with pytest.raises(AdapterError) as caught:
        parser(b"private arbitrary invalid payload", context)
    assert "private" not in str(caught.value)
