import json
from dataclasses import replace
from datetime import timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from bef_redbird.domain import (
    AnalysisContext,
    AnalysisSettings,
    DataError,
    QueueSettings,
    Review,
    ReviewState,
    select_queue,
)
from bef_redbird.export import dump_session, export_analysis, restore_session
from bef_redbird.service import build_analysis
from tests.fixtures import load_analysis


@pytest.fixture(scope="module")
def analysis():
    return load_analysis()


def test_full_filter_and_group_contract(analysis):
    assert analysis.snapshot.input_rows == 127
    assert analysis.snapshot.duplicate_rows == 1
    assert analysis.snapshot.conflicted_ids == {"900005"}
    assert analysis.filter_counts["Gyldige entydige produkter"] == 124
    assert analysis.filter_counts["Efter H/K-fravalg"] == 123
    assert analysis.filter_counts["Efter A/B → C"] == 122
    assert analysis.filter_counts["Efter kostpris < 1.000 kr."] == 121
    assert analysis.ranked_ids == tuple(f"{i:06d}" for i in range(1, 121))
    assert analysis.candidates["900004"].rank is None
    assert {v.product_id for v in analysis.group_members("000001")} == {
        "000001",
        "000002",
        "000003",
        "900001",
        "900002",
        "900003",
        "900004",
    }


@pytest.mark.parametrize("n", [1, 17, 50, 80, 100, 1000000000])
def test_queue_prefix_export_manifest(analysis, n):
    state = ReviewState(analysis.analysis_id, QueueSettings(top_n=n))
    expected = tuple(f"{i:06d}" for i in range(1, min(n, 120) + 1))
    assert select_queue(analysis, state.queue) == expected
    wb = load_workbook(BytesIO(export_analysis(analysis, state)), data_only=False)
    assert tuple(row[1].value for row in list(wb["Kandidater"].rows)[1:]) == expected
    metadata = list(wb["Metode og datakontrol"].values)
    assert tuple(v for k, v in metadata if k == "ordered_exported_id") == expected
    info = json.loads(next(v for k, v in metadata if k == "queue_export"))
    assert info == {
        "top_n": n,
        "mode": "top_n",
        "effective_n": min(n, 120),
        "scope": "queue",
        "exported_count": min(n, 120),
    }
    first = list(wb["Kandidater"].rows)[1]
    assert first[1].data_type == "s"
    assert first[3].value == 100
    assert first[7].value == 100
    assert first[3].data_type == "n"


@pytest.mark.parametrize("n", [0, -1, "17", 1.5, True, None, 1000000001])
def test_invalid_queue(n):
    with pytest.raises(DataError):
        QueueSettings(top_n=n)


def test_all_modes_small_and_empty():
    a = load_analysis(37)
    assert len(select_queue(a, QueueSettings(top_n=80))) == 37
    assert len(select_queue(a, QueueSettings("all_rankable", 1))) == 37
    empty = load_analysis(0)
    assert select_queue(empty, QueueSettings()) == ()
    with pytest.raises(DataError, match="Ingen produkter"):
        export_analysis(empty, ReviewState(empty.analysis_id))


def test_review_resize_export_scope_and_restore(analysis):
    state = ReviewState(
        analysis.analysis_id,
        QueueSettings(top_n=80),
        {"000080": Review("Interessant", "Gem denne vurdering")},
        {"000080"},
        "000080",
    )
    restored = restore_session(dump_session(analysis, state), analysis)
    restored.queue = QueueSettings(top_n=17)
    assert restored.reviews == state.reviews
    assert restored.selected_ids == {"000080"}
    assert restored.inspected_id == "000080"
    wb = load_workbook(BytesIO(export_analysis(analysis, restored, "selected")))
    rows = list(wb["Kandidater"].values)
    assert len(rows) == 2
    assert rows[1][1] == "000080"
    assert rows[1][25:27] == ("Interessant", "Gem denne vurdering")
    all_rows = list(
        load_workbook(BytesIO(export_analysis(analysis, restored, "all_eligible")))[
            "Kandidater"
        ].values
    )
    assert len(all_rows) == 122
    assert all_rows[-1][0] is None and all_rows[-1][1] == "900004"
    assert all_rows[-1][24]
    restored.queue = QueueSettings("all_rankable", 80)
    assert restore_session(dump_session(analysis, restored), analysis).queue == restored.queue


@pytest.mark.parametrize(
    "note", ["=1+1", "+SUM(A1)", "-1+2", "@A1", "\t=2", "\n=3", "https://example.invalid"]
)
def test_export_injection_is_plain_text(analysis, note):
    state = ReviewState(analysis.analysis_id, QueueSettings(top_n=1), {"000001": Review(note=note)})
    wb = load_workbook(BytesIO(export_analysis(analysis, state)), data_only=False)
    cell = wb["Kandidater"].cell(2, 27)
    assert cell.data_type == "s" and cell.value == note and cell.hyperlink is None


def test_identity_replay_settings_and_isolation(analysis):
    other = build_analysis(analysis.snapshot, analysis.history, analysis.context, analysis.settings)
    assert other.analysis_id == analysis.analysis_id
    assert other.run_id != analysis.run_id
    assert other.candidates == analysis.candidates
    reversed_snapshot = replace(
        analysis.snapshot, items=dict(reversed(list(analysis.snapshot.items.items())))
    )
    assert (
        build_analysis(reversed_snapshot, analysis.history, analysis.context).ranked_ids
        == analysis.ranked_ids
    )
    for context, settings, history in [
        (
            AnalysisContext(analysis.context.current_period - timedelta(days=14)),
            analysis.settings,
            analysis.history,
        ),
        (analysis.context, AnalysisSettings(25, 20), analysis.history),
        (
            analysis.context,
            analysis.settings,
            replace(analysis.history, source=replace(analysis.history.source, sha256="b" * 64)),
        ),
        (
            replace(analysis.context, cito_received_at="2025-01-01T12:00:00+01:00"),
            analysis.settings,
            analysis.history,
        ),
    ]:
        changed = build_analysis(analysis.snapshot, history, context, settings)
        assert changed.analysis_id != analysis.analysis_id
        with pytest.raises(DataError):
            restore_session(dump_session(analysis, ReviewState(analysis.analysis_id)), changed)
    one, two = ReviewState(analysis.analysis_id), ReviewState(analysis.analysis_id)
    one.reviews["000001"] = Review(note="Session A")
    assert not two.reviews


def test_stale_mismatch_missing_window(analysis):
    snapshot = replace(
        analysis.snapshot,
        items={"000001": replace(analysis.snapshot.items["000001"], current_cost_ore=99999)},
    )
    a = build_analysis(snapshot, analysis.history, analysis.context)
    assert a.candidates["000001"].flags == ("cost_mismatch",)
    assert a.candidates["000001"].features.current_aip_ore == 10000
    assert a.candidates["000001"].item.current_cost_ore == 99999
    stale = build_analysis(
        analysis.snapshot,
        analysis.history,
        AnalysisContext(analysis.context.current_period + timedelta(days=14)),
    )
    assert not stale.ranked_ids
    assert "unaligned_period" in stale.candidates["000001"].flags
    assert any(issue.code == "older_period_match" for issue in stale.issues)
    assert stale.context.current_period == analysis.context.current_period + timedelta(days=14)
    short = load_analysis(2, 4)
    assert not short.ranked_ids
    assert short.candidates["000001"].features.coverage == 1
    assert short.candidates["000001"].features.window_periods == 4


def test_session_rejects_malformed_unknown_and_oversized(analysis):
    for data in (b"not json", b"[]", b"x" * (2 * 1024 * 1024 + 1)):
        with pytest.raises(DataError):
            restore_session(data, analysis)
    payload = json.loads(dump_session(analysis, ReviewState(analysis.analysis_id)))
    payload["selected_ids"] = ["123456"]
    with pytest.raises(DataError):
        restore_session(json.dumps(payload).encode(), analysis)
