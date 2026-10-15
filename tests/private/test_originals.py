"""Opt-in original-file checks. Output is aggregate-only; originals stay immutable."""

import hashlib
import json
import os
import time
from datetime import date
from decimal import Decimal
from pathlib import Path

import psutil
import pytest
import xlrd
from openpyxl import load_workbook

from bef_redbird.domain import AnalysisContext
from bef_redbird.imports import parse_cito, parse_lms_history
from bef_redbird.service import build_analysis

pytestmark = pytest.mark.private


def test_authorised_originals():
    paths = [os.environ.get(k) for k in ("REDBIRD_PRIVATE_CITO_PATH", "REDBIRD_PRIVATE_LMS_PATH")]
    if not all(paths):
        pytest.skip("Private paths not supplied; original-file acceptance not executed.")
    cito_path, lms_path = (Path(p) for p in paths)
    assert cito_path.is_absolute() and lms_path.is_absolute()
    original_hashes = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (cito_path, lms_path)]
    started = time.perf_counter()
    cito = parse_cito(cito_path.read_bytes(), cito_path.name)
    history = parse_lms_history(lms_path.read_bytes(), lms_path.name)
    import_seconds = time.perf_counter() - started
    # This is a declared test context, not an inferred report-date assertion in the UI.
    current = date.fromisoformat(os.environ.get("REDBIRD_PRIVATE_PERIOD", "2026-09-28"))
    result = build_analysis(cito, history, AnalysisContext(current))

    # Separate raw-cell Cito check without calling production normalisation/filter helpers.
    if cito_path.suffix.lower() == ".xls":
        workbook = xlrd.open_workbook(str(cito_path), on_demand=True)
        sheet = workbook.sheet_by_name(cito.sheet_name)
        raw_rows = [sheet.row_values(i) for i in range(sheet.nrows)]
        workbook.release_resources()
    else:
        workbook_x = load_workbook(cito_path, read_only=True, data_only=False)
        worksheet = workbook_x["Table"]
        worksheet.reset_dimensions()
        raw_rows = list(worksheet.values)
        workbook_x.close()
    headers = [str(v).strip().lower().replace("ny subkode", "nysubkode") for v in raw_rows[0]]
    indices = [
        headers.index(k)
        for k in ("varenr", "opbevkode", "kostpris", "subkode", "nysubkode", "subgrp")
    ]
    unique = set()
    raw_count = 0
    for row in raw_rows[1:]:
        values = [row[i] for i in indices]
        if values[0] in (None, ""):
            continue
        raw_count += 1
        pid = f"{int(Decimal(str(values[0]))):06d}"
        storage = str(values[1] or "").strip().upper()
        price = int((Decimal(str(values[2])) * 100).to_integral_value())
        group = str(int(Decimal(str(values[5])))) if values[5] not in (None, "") else None
        unique.add(
            (
                pid,
                storage,
                price,
                str(values[3] or "").strip().upper(),
                str(values[4] or "").strip().upper(),
                group,
            )
        )
    after_storage = [r for r in unique if r[1] not in {"H", "K"}]
    after_codes = [r for r in after_storage if r[3] in {"A", "B"} and r[4] == "C"]
    eligible = [r for r in after_codes if r[2] < 100000]
    assert raw_count == cito.input_rows
    assert len(unique) == cito.unique_required_rows
    assert len(after_storage) == result.filter_counts["Efter H/K-fravalg"]
    assert len(after_codes) == result.filter_counts["Efter A/B → C"]
    assert {r[0] for r in eligible} == set(result.candidates)

    # Independent streaming count and selected source cells, never printed or persisted.
    wb = load_workbook(lms_path, read_only=True, data_only=False, keep_links=False)
    ws = wb["Medicinpriser data"]
    ws.reset_dimensions()
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    pid_col, indicator_col = header.index("Varenummer"), header.index("Indikator")
    date_cols = [i for i, v in enumerate(header) if str(v).isdigit() and len(str(v)) == 8]
    observed, ids, sampled = 0, set(), 0
    for row in rows:
        row = (*row, *([None] * (len(header) - len(row))))
        if row[indicator_col] != "AIP":
            continue
        pid = f"{int(Decimal(str(row[pid_col]))):06d}"
        ids.add(pid)
        observed += sum(row[col] not in (None, "") for col in date_cols)
        if pid in result.candidates and sampled < 12:
            expected = tuple(
                None
                if row[c] in (None, "")
                else int((Decimal(str(row[c])) * 100).to_integral_value())
                for c in date_cols
            )
            assert history.prices[pid] == expected
            sampled += 1
    wb.close()
    assert ids == set(history.prices)
    assert observed == sum(v is not None for series in history.prices.values() for v in series)
    assert sampled == min(12, len(set(result.candidates) & ids))
    summary = {
        "hashes": original_hashes,
        "input_rows": raw_count,
        "unique_rows": len(unique),
        "after_storage": len(after_storage),
        "after_codes": len(after_codes),
        "eligible": len(eligible),
        "candidate_groups": result.filter_counts["Kandidatgrupper"],
        "aip_ids": len(ids),
        "periods": len(date_cols),
        "aip_observations": observed,
        "candidate_history": len(set(result.candidates) & ids),
        "matched_current": sum(
            c.features.current_aip_ore == c.item.current_cost_ore
            for c in result.candidates.values()
        ),
        "current_mismatches": sum("cost_mismatch" in c.flags for c in result.candidates.values()),
        "rankable": len(result.ranked_ids),
        "issues": len(result.issues),
        "import_seconds": round(import_seconds, 2),
        "peak_process_mib": round(
            getattr(psutil.Process().memory_info(), "peak_wset", psutil.Process().memory_info().rss)
            / 1024**2,
            1,
        ),
        "source_samples_checked": sampled,
    }
    if original_hashes == [
        # pragma: allowlist nextline secret
        "f30c17c0b85109630ac5a53d8c0dac0073dfd8cba8ea91d17d3c2eb1c47bcfcf",
        # pragma: allowlist nextline secret
        "9ee23e8a83a1d27f1cf60eafc4d2274406ab49f502a7b02238a2e30d1cb04dd8",
    ]:
        assert [
            raw_count,
            len(unique),
            len(after_storage),
            len(after_codes),
            len(eligible),
            len(ids),
            len(date_cols),
            observed,
        ] == [40954, 6828, 6205, 509, 479, 16991, 135, 1442961]
    assert original_hashes == [
        hashlib.sha256(p.read_bytes()).hexdigest() for p in (cito_path, lms_path)
    ]
    print(json.dumps(summary, sort_keys=True))
