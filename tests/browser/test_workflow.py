"""Real Chromium proof with wholly invented workbooks and isolated sessions."""

from __future__ import annotations

import json
import os
import re
import socket
import subprocess
import sys
import time
from collections.abc import Iterator
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING, Any
from urllib.error import URLError
from urllib.request import urlopen

import pytest
from openpyxl import load_workbook

from tests.fixtures import synthetic_workbooks

if TYPE_CHECKING:
    from playwright.sync_api import Browser, Locator, Page
playwright = pytest.importorskip(
    "playwright.sync_api", reason="Install the browser dependency group"
)
expect = playwright.expect
pytestmark = pytest.mark.browser
ROOT = Path(__file__).resolve().parents[2]
LONG_NAME = (
    "Eksempellægemiddel med langt navn 100 mg filmovertrukne tabletter · 100 stk. blisterpakning"
)


@pytest.fixture(scope="module")
def workspace_url() -> Iterator[str]:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
    process = subprocess.Popen(
        [
            sys.executable,
            "-m",
            "streamlit",
            "run",
            "app.py",
            "--server.address=127.0.0.1",
            f"--server.port={port}",
            "--server.headless=true",
            "--browser.gatherUsageStats=false",
        ],
        cwd=ROOT,
        env={**os.environ, "PYTHONUTF8": "1"},
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    address = f"http://127.0.0.1:{port}"
    try:
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline:
            if process.poll() is not None:
                pytest.fail("Synthetic Streamlit test server stopped before becoming ready.")
            try:
                with urlopen(f"{address}/_stcore/health", timeout=1) as response:
                    if response.status == 200:
                        break
            except (URLError, TimeoutError):
                time.sleep(0.1)
        else:
            pytest.fail("Synthetic Streamlit test server did not become ready.")
        yield address
    finally:
        process.terminate()
        try:
            process.wait(timeout=10)
        except subprocess.TimeoutExpired:
            process.kill()
            process.wait(timeout=10)


def _settled(page: Page) -> None:
    expect(page.get_by_test_id("stApp")).to_have_attribute("data-test-script-state", "notRunning")
    # Widget values update before queued server reruns and streamed DOM updates.
    # Require actual render quiescence before interacting with the next control.
    page.wait_for_function("""() => new Promise(resolve => {
        const app = document.querySelector('[data-testid="stApp"]');
        let timer;
        const observer = new MutationObserver(check);
        function check() {
            clearTimeout(timer);
            if (app.getAttribute('data-test-script-state') === 'notRunning') {
                timer = setTimeout(() => {observer.disconnect(); resolve(true)}, 200);
            }
        }
        observer.observe(app, {attributes:true, childList:true, characterData:true, subtree:true});
        check();
    })""")


def _expect_product(page: Page, product_id: str) -> None:
    expect(
        page.locator(".st-key-price_workspace")
        .get_by_test_id("stText")
        .filter(has_text=re.compile(rf"^{product_id} · "))
    ).to_be_visible()


def _expect_table_product(page: Page, product_id: str | None) -> None:
    """Read the native rendered accessibility rows; never assign browser selection."""
    grid = page.locator(".st-key-candidate_list").get_by_test_id("stDataFrame")
    selected = grid.locator('tbody tr[aria-selected="true"]')
    if product_id is None:
        expect(selected).to_have_count(0)
        expect(grid.get_by_role("button", name="Clear selection", exact=True)).to_have_count(0)
    else:
        expect(selected).to_have_count(1)
        expect(selected.locator('[aria-colindex="2"]')).to_have_text(product_id)
        expect(grid.get_by_role("button", name="Clear selection", exact=True)).to_have_count(0)


def _open_from_table(page: Page, product_id: str, *, query: str | None = None) -> None:
    """Filter and click the real native grid, then restore the user's search."""
    page.keyboard.press("Escape")
    search = page.get_by_role("textbox", name="Søg i kandidatlisten", exact=True)
    previous = search.input_value()
    search.fill(query or product_id)
    search.press("Enter")
    _settled(page)
    expect(page.get_by_text("1 synlige efter søgning og filtre.", exact=False)).to_be_visible()
    grid = page.locator(".st-key-candidate_list").get_by_test_id("stDataFrame")
    grid.evaluate("""e => {
        e.scrollIntoView({block:'start'});
        e.closest('[data-testid="stMain"]').scrollBy(0, -80);
    }""")
    expect(grid.locator("canvas").first).to_be_visible()
    grid.click(position={"x": 16, "y": 56})
    _settled(page)
    _expect_product(page, product_id)
    _expect_table_product(page, product_id)
    expect(page.get_by_role("button", name="Forrige kandidat", exact=True)).to_be_disabled()
    expect(page.get_by_role("button", name="Næste kandidat", exact=True)).to_be_disabled()
    search.fill(previous)
    search.press("Enter")
    _settled(page)
    _expect_product(page, product_id)


def _alternative(page: Page, product_id: str) -> Locator:
    page.keyboard.press("Escape")
    _settled(page)
    checkbox = page.get_by_role("checkbox", name=product_id, exact=True)
    checkbox.evaluate("e => e.scrollIntoView({block:'center', behavior:'instant'})")
    return checkbox


def _review_alternative(page: Page, product_id: str) -> None:
    checkbox = _alternative(page, product_id)
    row = checkbox.locator("xpath=ancestor::*[contains(@class, 'st-key-alternative_')][1]")
    row.get_by_role("button", name="Vis", exact=True).click()
    _settled(page)


def _toggle_alternative(page: Page, product_id: str, *, checked: bool = True) -> None:
    checkbox = _alternative(page, product_id)
    expect(checkbox).to_be_enabled()
    # Native Streamlit's checkbox input is covered by its associated label.
    # Exercise the accessible keyboard interaction without bypassing event handlers.
    if checkbox.is_checked() != checked:
        checkbox.press("Space")
    expect(checkbox).to_be_checked(checked=checked)
    _settled(page)


def _select(page: Page, label: str, value: str) -> None:
    page.keyboard.press("Escape")
    _settled(page)
    control = page.get_by_role("combobox", name=label, exact=True)
    control.evaluate("e => e.scrollIntoView({block:'center', behavior:'instant'})")
    expect(control).to_be_in_viewport(ratio=1)
    control.click()
    page.get_by_role("option", name=value, exact=True).click()
    expect(control).to_have_value(value)
    _settled(page)


def _load(
    page: Page,
    workspace_url: str,
    tmp_path: Path,
    count: int = 120,
    *,
    descriptive_labels: bool = False,
    large_group: bool = False,
    varied_prices: bool = False,
) -> None:
    cito, history, period = synthetic_workbooks(count=count, periods=40)
    if large_group or varied_prices:
        workbook = load_workbook(BytesIO(cito))
        for row in workbook["Table"].iter_rows(min_row=2):
            # Keep the deliberately conflicting fixture outside this group;
            # one-screen measurements exercise a healthy comparison.
            if large_group and row[0].value != "900005":
                row[5].value = "1"
            if varied_prices and 0 < int(row[0].value) <= count:
                row[2].value += (int(row[0].value) % 7) * 5
        output = BytesIO()
        workbook.save(output)
        cito = output.getvalue()
    if descriptive_labels or varied_prices:
        # Display-only metadata exercises ordinary long pharmacy descriptions.
        workbook = load_workbook(BytesIO(history))
        sheet = workbook["Medicinpriser data"]
        company_col = sheet.max_column + 1
        if descriptive_labels:
            sheet.cell(1, company_col, "Firma")
        for row in sheet.iter_rows(min_row=2):
            if row[1].value == "AIP":
                if descriptive_labels:
                    row[2].value = LONG_NAME
                    sheet.cell(row[0].row, company_col, f"Firma {row[0].value}")
                if varied_prices and 0 < int(row[0].value) <= count:
                    for cell in row[3 : company_col - 1]:
                        cell.value += (int(row[0].value) % 7) * 5
        output = BytesIO()
        workbook.save(output)
        history = output.getvalue()
    page.goto(workspace_url)
    # The server health endpoint can be ready before a cold app import has finished.
    expect(page.get_by_text("Upload Cito-rapport og LMS-prishistorik", exact=False)).to_be_visible(
        timeout=30000
    )
    page.screenshot(path=str(tmp_path / "empty-laptop.png"))
    uploads = page.locator('input[type="file"]')
    uploads.nth(0).set_input_files(
        {
            "name": "synthetic-cito.xlsx",
            "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "buffer": cito,
        }
    )
    expect(
        page.get_by_role("button", name="Remove synthetic-cito.xlsx", exact=True)
    ).to_be_visible()
    _settled(page)
    uploads.nth(1).set_input_files(
        {
            "name": "synthetic-lms.xlsx",
            "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "buffer": history,
        }
    )
    expect(page.get_by_role("button", name="Remove synthetic-lms.xlsx", exact=True)).to_be_visible()
    _settled(page)
    date_input = page.get_by_test_id("stDateInput").locator("input")
    date_input.fill(period.isoformat())
    date_input.press("Enter")
    _settled(page)
    page.get_by_role("button", name="Indlæs og kontroller", exact=True).click()
    expect(
        page.get_by_text(f"Viser {min(50, count)} af 50 ønskede kandidater", exact=False)
    ).to_be_visible(timeout=20000)
    page.get_by_role("heading", name="Redbird · indkøbsgennemgang", exact=True).evaluate(
        "e => e.scrollIntoView({block:'start'})"
    )
    expect(
        page.get_by_text(
            "Prioritering til gennemgang — ikke dokumenteret købsgevinst.", exact=False
        )
    ).to_be_in_viewport(ratio=1)
    expect(page.get_by_role("switch", name="Vis vurdering", exact=True)).not_to_be_checked()
    expect(page.locator(".st-key-review_editor")).to_have_count(0)


def _show_review(page: Page, *, enabled: bool = True) -> None:
    toggle = page.get_by_role("switch", name="Vis vurdering", exact=True)
    if toggle.is_checked() != enabled:
        toggle.press("Space")
        _settled(page)
    expect(toggle).to_be_checked(checked=enabled)
    expect(page.locator(".st-key-review_editor")).to_have_count(1 if enabled else 0)


def _download(page: Page, label: str, tmp_path: Path) -> Path:
    with page.expect_download() as event:
        page.get_by_role("button", name=label, exact=True).click()
    result = event.value
    target = tmp_path / result.suggested_filename
    result.save_as(target)
    return target


def _visible_height(locator: Locator) -> float:
    """Measure visible height after viewport and ancestor clipping."""
    return float(
        locator.evaluate("""element => {
        const rect = element.getBoundingClientRect();
        let top = Math.max(64, rect.top), bottom = Math.min(innerHeight, rect.bottom);
        for (let parent = element.parentElement; parent; parent = parent.parentElement) {
            if (/(auto|scroll|hidden)/.test(getComputedStyle(parent).overflowY)) {
                const clip = parent.getBoundingClientRect();
                top = Math.max(top, clip.top); bottom = Math.min(bottom, clip.bottom);
            }
        }
        return Math.max(0, bottom - top);
    }""")
    )


def _linked_layout_evidence(page: Page, tmp_path: Path) -> dict[str, Any]:
    page.keyboard.press("Escape")
    _settled(page)
    queue = page.locator(".st-key-candidate_list")
    prices = page.locator(".st-key-price_workspace")
    form = page.locator(".st-key-review_editor")
    expect(page.get_by_text("Vælg fra listen med tastatur", exact=True)).to_have_count(0)
    expect(
        page.get_by_role("combobox", name="Kandidat i den viste liste", exact=True)
    ).to_have_count(0)
    grid = queue.get_by_test_id("stDataFrame")
    chart = prices.get_by_test_id("stPlotlyChart").first
    price_metrics = prices.get_by_test_id("stMetric")
    for label in ("Forrige kandidat", "Næste kandidat"):
        expect(prices.get_by_role("button", name=label, exact=True)).to_have_attribute(
            "kind", "secondary"
        )
    rows = prices.locator('[class*="st-key-alternative_"]')
    row_count = rows.count()
    assert row_count in (2, 6, 12)
    expect(
        prices.get_by_test_id("stMarkdownContainer").filter(
            has_text=re.compile(rf"^Substitutionsalternativer: {row_count} · \d valgt\s*$")
        )
    ).to_be_visible()
    for choice in ("Alle", "26 perioder", "52 perioder"):
        expect(prices.get_by_role("radio", name=choice, exact=True)).to_be_visible()
    expect(prices.get_by_role("checkbox", name="Senere priser", exact=True)).to_be_visible()
    expect(page.get_by_role("button", name="Flere grafvalg", exact=True)).to_have_count(0)
    for link in ("Se den valgte kandidats priser", "Gem din vurdering", "Til kandidatlisten"):
        expect(page.get_by_role("link", name=link, exact=True)).to_have_count(0)
    expect(prices.get_by_text(re.compile(r"^Vis flere alternativer"))).to_have_count(0)
    for row in rows.all():
        # Every alternative is a direct, visible row with aligned native fields.
        expect(row).to_be_visible()
        expect(row.get_by_role("checkbox")).to_have_count(1)
        expect(row.get_by_test_id("stText")).to_have_count(2)
        expect(row.locator("[data-testid='stMarkdownContainer'] strong")).to_have_count(1)
    expect(price_metrics.get_by_test_id("stMetricLabel")).to_have_text(
        ["Kostpris fra LMS", "Median i egen historik", "Afstand til median"]
    )
    documentation = page.get_by_text("Kilder, metode og datakontrol", exact=True)
    metrics = {}
    expect(page.get_by_role("tab")).to_have_count(0)
    for size, width, height in (
        ("laptop", 1366, 768),
        ("desktop", 1920, 1080),
        ("narrow", 600, 900),
    ):
        page.set_viewport_size({"width": width, "height": height})
        # Plotly resizes asynchronously after the native container has reflowed.
        page.wait_for_function("""() => {
            const plot = document.querySelector('.js-plotly-plot');
            return plot && Math.abs(plot._fullLayout.width - plot.clientWidth) < 3;
        }""")
        prices.screenshot(path=str(tmp_path / f"measured-prices-{size}.png"))
        sections = [section.bounding_box() for section in (queue, prices)]
        review_box = form.bounding_box() if form.count() else None
        if review_box is not None:
            sections.append(review_box)
        (tmp_path / f"section-dimensions-{size}.json").write_text(
            json.dumps(sections, indent=2), encoding="utf-8"
        )
        assert all(sections)
        candidate_box, price_box = sections[:2]
        assert candidate_box is not None and price_box is not None
        if width > 600:
            price_cards = [card.bounding_box() for card in price_metrics.all()]
            assert all(price_cards)
            first, middle, last = price_cards
            assert first is not None and middle is not None and last is not None
            assert first["x"] + first["width"] < middle["x"]
            assert middle["x"] + middle["width"] < last["x"]
            assert max(card["y"] for card in (first, middle, last)) - first["y"] <= 2
            assert abs(candidate_box["y"] - price_box["y"]) <= 2, sections
            assert abs(candidate_box["height"] - price_box["height"]) <= 2, sections
            assert candidate_box["x"] + candidate_box["width"] < price_box["x"], sections
            assert price_box["width"] > candidate_box["width"], sections
            assert candidate_box["width"] >= 450, sections
            if review_box is not None:
                assert review_box["width"] > candidate_box["width"] + price_box["width"], sections
        else:
            assert candidate_box["y"] + candidate_box["height"] <= price_box["y"], sections
            assert abs(candidate_box["x"] - price_box["x"]) <= 2, sections
            assert abs(candidate_box["width"] - price_box["width"]) <= 2, sections
            assert candidate_box["width"] >= 450, sections
        comparison_bottom = max(
            candidate_box["y"] + candidate_box["height"],
            price_box["y"] + price_box["height"],
        )
        if review_box is not None:
            assert review_box["y"] >= comparison_bottom, sections
            assert abs(review_box["x"] - candidate_box["x"]) <= 2, sections
            comparison_bottom = review_box["y"] + review_box["height"]
        grid_box = grid.bounding_box()
        (tmp_path / f"grid-dimensions-{size}.json").write_text(
            json.dumps(grid_box, indent=2), encoding="utf-8"
        )
        assert grid_box is not None and grid_box["width"] >= candidate_box["width"] - 40
        assert 419 <= grid_box["height"] <= 421, grid_box
        chart_box = chart.bounding_box()
        assert chart_box is not None
        assert 359 <= chart_box["height"] <= 361, chart_box
        dimensions = chart.locator(".js-plotly-plot").evaluate("""plot => ({
            width: plot._fullLayout._size.w,
            height: plot._fullLayout._size.h,
            font: plot._fullLayout.font.size,
        })""")
        assert 243 <= dimensions["height"] <= 245, dimensions
        assert dimensions["width"] >= (500 if width > 600 else 300), dimensions
        assert dimensions["font"] >= 12, dimensions
        trace_count = chart.locator(".js-plotly-plot").evaluate("plot => plot.data.length")
        legend_labels = chart.locator(".legendtext")
        expect(legend_labels).to_have_count(trace_count)
        expect(chart.locator(".legend .scrollbar")).to_be_hidden()
        method_box = page.get_by_role(
            "heading", name="Sådan læses analysen", exact=True
        ).bounding_box()
        assert method_box is not None
        assert method_box["y"] > comparison_bottom
        assert method_box["width"] <= min(1100, width - 32) + 1, method_box
        documentation_box = documentation.bounding_box()
        assert documentation_box is not None
        assert documentation_box["y"] > comparison_bottom
        overflow = page.evaluate("""() => ({
            viewport: innerWidth,
            page: document.documentElement.scrollWidth,
            controls: [...document.querySelectorAll('input,textarea,button')]
                .filter(element => {
                    const rect = element.getBoundingClientRect();
                    return element.checkVisibility({checkOpacity:true, checkVisibilityCSS:true})
                        && rect.width && rect.height
                        && (rect.left < -1 || rect.right > innerWidth+1);
                }).map(element => element.getAttribute('aria-label') || element.textContent),
        })""")
        assert overflow["page"] <= overflow["viewport"] + 1, overflow
        assert not overflow["controls"], overflow
        row_layout = rows.evaluate_all("""elements => elements.map(row => ({
            id:row.querySelector('input').getAttribute('aria-label'),
            bounds:(() => {const r=row.getBoundingClientRect();
                return {y:r.y,height:r.height};})(),
            hiddenDisclosure:!!row.closest('details'),
            nestedScroll:[...function*(node){
                for (let p=node.parentElement;p && !p.matches('[data-testid="stMain"]');
                     p=p.parentElement) yield p;
            }(row)].some(p => /(auto|scroll)/.test(getComputedStyle(p).overflowY)
                && p.scrollHeight > p.clientHeight+2),
            columns:[...row.querySelectorAll('[data-testid="stColumn"]')].map(column => {
                const rect=column.getBoundingClientRect();
                return {x:rect.x,y:rect.y,width:rect.width,height:rect.height};
            }),
        }))""")
        assert all(not row["hiddenDisclosure"] and not row["nestedScroll"] for row in row_layout)
        assert row_layout[0]["bounds"]["y"] >= chart_box["y"] + chart_box["height"]
        assert all(
            first["bounds"]["y"] + first["bounds"]["height"] <= second["bounds"]["y"] + 1
            for first, second in zip(row_layout, row_layout[1:], strict=False)
        ), row_layout
        assert all(
            second["bounds"]["y"] - first["bounds"]["y"] - first["bounds"]["height"] <= 24
            for first, second in zip(row_layout, row_layout[1:], strict=False)
        ), row_layout
        last_row = row_layout[-1]["bounds"]
        assert last_row["y"] + last_row["height"] <= price_box["y"] + price_box["height"]
        if width > 600:
            for row in row_layout:
                columns = row["columns"]
                assert len(columns) == 5, row
                assert all(
                    left["x"] + left["width"] <= right["x"]
                    for left, right in zip(columns, columns[1:], strict=False)
                ), row
        grid_scroll = grid.evaluate("""element => [...element.querySelectorAll('*')]
            .filter(node => /(auto|scroll)/.test(getComputedStyle(node).overflowX))
            .map(node => ({width:node.clientWidth, content:node.scrollWidth}))""")
        assert grid_scroll, "The native table scroll viewport must be measurable"
        assert all(item["content"] <= item["width"] + 2 for item in grid_scroll), grid_scroll
        expect(grid.locator("th")).to_have_text(
            ["Prioritet", "Varenr.", "Lægemiddel, styrke og pakning", "Kostpris, kr."]
        )
        metrics[size] = {
            "sections": sections,
            "grid": grid_box,
            "plot": dimensions,
            "chart": chart_box,
            "grid_scroll": grid_scroll,
            "alternative_count": row_count,
            "alternative_layout": row_layout,
        }
        (tmp_path / "layout-metrics.json").write_text(
            json.dumps(metrics, indent=2), encoding="utf-8"
        )
        if width > 600:
            queue.evaluate("""e => {
                e.scrollIntoView({block:'start'});
                e.closest('[data-testid="stMain"]').scrollBy(0, -64);
            }""")
            expect(chart).to_be_in_viewport(ratio=1)
            for label in chart.locator(".legendtext").all():
                expect(label).to_be_in_viewport(ratio=1)
            expect(
                prices.get_by_role("button", name="Gruppedetaljer", exact=True)
            ).to_be_in_viewport(ratio=1)
            for choice in ("Alle", "26 perioder", "52 perioder"):
                expect(prices.get_by_text(choice, exact=True)).to_be_in_viewport(ratio=1)
            expect(prices.get_by_text("Senere priser", exact=True)).to_be_in_viewport(ratio=1)
            page.screenshot(path=str(tmp_path / f"linked-whole-workspace-{size}.png"))
        else:
            for position, row in (("top", rows.first), ("bottom", rows.last)):
                row.evaluate("""e => {
                    e.scrollIntoView({block:'start'});
                    e.closest('[data-testid="stMain"]').scrollBy(0, -80);
                }""")
                expect(row).to_be_in_viewport(ratio=1)
                checkbox = row.get_by_role("checkbox")
                product_id = checkbox.get_attribute("aria-label")
                assert product_id is not None and re.fullmatch(r"\d{6}", product_id)
                expect(checkbox).to_have_accessible_name(product_id)
                label = checkbox.locator("xpath=ancestor::label[1]")
                expect(label).to_be_in_viewport(ratio=1)
                expect(label.get_by_text(product_id, exact=True)).to_be_in_viewport(ratio=1)
                page.screenshot(path=str(tmp_path / f"linked-alternatives-narrow-{position}.png"))
        for section, target in (
            ("controls", queue),
            ("table", grid),
            ("prices", prices),
            ("chart", chart),
            *((("review", form),) if form.count() else ()),
        ):
            target.evaluate("""e => {
                e.scrollIntoView({block:'start'});
                e.closest('[data-testid="stMain"]').scrollBy(0, -80);
            }""")
            if section == "chart":
                assert _visible_height(chart) >= 258
                for label in legend_labels.all():
                    expect(label).to_be_in_viewport(ratio=1)
            page.screenshot(path=str(tmp_path / f"linked-{section}-{size}.png"))
    (tmp_path / "layout-metrics.json").write_text(json.dumps(metrics, indent=2), encoding="utf-8")
    page.set_viewport_size({"width": 1366, "height": 768})
    return metrics


def _unified_hover_evidence(page: Page, tmp_path: Path) -> None:
    chart = page.get_by_test_id("stPlotlyChart").first
    chart.evaluate("e => e.scrollIntoView({block:'center'})")
    plot = chart.locator(".js-plotly-plot")
    evidence = plot.evaluate("""plot => ({
        ids:plot.data.map(trace => trace.meta),
        values:plot.data.map(trace => trace.y),
        periods:plot.data.map(trace => trace.x),
        colors:plot.data.map(trace => trace.line.color),
        dashes:plot.data.map(trace => trace.line.dash),
        symbols:plot.data.map(trace => trace.marker.symbol),
        hovermode:plot._fullLayout.hovermode,
    })""")
    assert evidence["ids"] == [f"{i:06d}" for i in range(1, 7)]
    assert all(values == evidence["values"][0] for values in evidence["values"])
    for style in ("colors", "dashes", "symbols"):
        assert len(set(evidence[style])) == 6, evidence
    assert evidence["hovermode"] == "x unified", evidence
    # Move the real pointer onto an existing observation; never move or jitter source points.
    point = plot.evaluate("""plot => {
        const trace=plot.data[0], index=20, layout=plot._fullLayout;
        const rect=plot.getBoundingClientRect();
        return {
            x:rect.x+layout.xaxis._offset+layout.xaxis.d2p(trace.x[index]),
            y:rect.y+layout.yaxis._offset+layout.yaxis.d2p(trace.y[index]),
        };
    }""")
    page.mouse.move(point["x"], point["y"])
    hover = chart.locator(".hoverlayer")
    for product_id in evidence["ids"]:
        expect(hover).to_contain_text(product_id)
    evidence["hover_box"] = hover.bounding_box()
    chart.screenshot(path=str(tmp_path / "six-identical-prices-unified-hover.png"))
    (tmp_path / "unified-hover-metrics.json").write_text(
        json.dumps(evidence, indent=2), encoding="utf-8"
    )
    hover_box = evidence["hover_box"]
    viewport = page.viewport_size
    assert hover_box is not None and viewport is not None
    # SVG intersection ratios can round below 1 for fully visible content.
    assert -1 <= hover_box["x"] <= viewport["width"]
    assert -1 <= hover_box["y"] <= viewport["height"]
    assert hover_box["x"] + hover_box["width"] <= viewport["width"] + 1
    assert hover_box["y"] + hover_box["height"] <= viewport["height"] + 1
    assert plot.evaluate("plot => plot.data.map(trace => trace.y)") == evidence["values"]
    assert plot.evaluate("plot => plot.data.map(trace => trace.x)") == evidence["periods"]
    evidence["hover_text"] = hover.text_content()
    chart.screenshot(path=str(tmp_path / "six-identical-prices-unified-hover.png"))
    (tmp_path / "unified-hover-metrics.json").write_text(
        json.dumps(evidence, indent=2), encoding="utf-8"
    )
    page.mouse.move(0, 0)


def _group_details_evidence(page: Page, tmp_path: Path, count: int = 7) -> None:
    evidence = {}
    for size, width, height in (
        ("laptop", 1366, 768),
        ("desktop", 1920, 1080),
        ("narrow", 600, 900),
    ):
        page.set_viewport_size({"width": width, "height": height})
        page.get_by_role("button", name="Gruppedetaljer", exact=True).click()
        body = page.get_by_test_id("stPopoverBody")
        cards = body.locator('[class*="st-key-group_detail_"]')
        expect(cards).to_have_count(count)
        expect(body.get_by_test_id("stDataFrame")).to_have_count(0)
        for card in cards.all():
            expect(card).to_be_visible()
            fields = card.get_by_test_id("stText")
            expect(fields).to_have_count(5)
            pid = re.search(r"\d{6}", card.inner_text())
            assert pid is not None
            product_id = pid.group()
            if product_id not in ("900002", "900003", "900004"):
                expect(fields.nth(0)).to_have_text(f"{LONG_NAME} · Firma {product_id}")
                expect(fields.nth(1)).to_have_text(f"Firma: Firma {product_id}")
            expect(fields.nth(2)).to_contain_text("Cito-kostpris:")
            expect(fields.nth(3)).to_contain_text("Opbevaring:")
            expect(fields.nth(4)).to_contain_text("Historik:")
        overflow = body.evaluate("""element => ({
            width: element.clientWidth, content: element.scrollWidth,
            fields: [...element.querySelectorAll('[data-testid="stText"]')].map(node => ({
                width:node.clientWidth, content:node.scrollWidth,
                height:node.clientHeight, contentHeight:node.scrollHeight,
                text:node.textContent,
            })),
        })""")
        assert overflow["content"] <= overflow["width"] + 1, overflow
        assert all(
            item["content"] <= item["width"] + 1 and item["contentHeight"] <= item["height"] + 1
            for item in overflow["fields"]
        ), overflow
        evidence[size] = overflow
        page.screenshot(path=str(tmp_path / f"group-details-{size}.png"))
        page.keyboard.press("Escape")
    (tmp_path / "group-details-metrics.json").write_text(
        json.dumps(evidence, indent=2), encoding="utf-8"
    )
    page.set_viewport_size({"width": 1366, "height": 768})


def test_descriptive_layout(browser: Browser, workspace_url: str, tmp_path: Path) -> None:
    context = browser.new_context(viewport={"width": 1366, "height": 768})
    page = context.new_page()
    page.set_default_timeout(10000)
    try:
        _load(page, workspace_url, tmp_path, descriptive_labels=True)
        expect(
            page.get_by_test_id("stText").filter(
                has_text=re.compile(r"^000001 · Eksempellægemiddel")
            )
        ).to_be_visible()
        _group_details_evidence(page, tmp_path)
        for product_id in ("000002", "000003", "900001"):
            _toggle_alternative(page, product_id)
        expect(
            page.locator(".st-key-price_workspace").get_by_role("button", name="Vis", exact=True)
        ).to_have_count(3)
        six_rows = _linked_layout_evidence(page, tmp_path)
        expect(page.locator(".st-key-review_editor")).to_have_count(0)
        # Group 3 is healthy; group 2 intentionally contains a conflicting fixture row.
        _open_from_table(page, "000007")
        small_group = tmp_path / "small-group"
        small_group.mkdir()
        two_rows = _linked_layout_evidence(page, small_group)
        for size in ("laptop", "desktop", "narrow"):
            assert six_rows[size]["chart"]["height"] == two_rows[size]["chart"]["height"]
            assert six_rows[size]["sections"][1]["height"] > two_rows[size]["sections"][1]["height"]
        expect(page.locator(".st-key-review_editor")).to_have_count(0)
        _open_from_table(page, "000004")
        expect(page.get_by_text("Der er uafklarede rækker i gruppen.", exact=False)).to_be_visible()
        page.locator(".st-key-price_workspace").screenshot(
            path=str(tmp_path / "incomplete-group-warning.png")
        )
        for size, width, height in (("laptop", 1366, 768), ("desktop", 1920, 1080)):
            page.set_viewport_size({"width": width, "height": height})
            for section, label in (
                ("flow", "Fra Cito-kandidat til din vurdering"),
                ("settings", "Anvendte indstillinger og tidsgrænse"),
                ("calculations", "Beregninger og præcise grænser"),
                ("sources", "Kilder, metode og datakontrol"),
            ):
                documentation = page.get_by_text(label, exact=True)
                documentation.click()
                _settled(page)
                # Native details expansion animates independently of Streamlit reruns.
                documentation.locator("xpath=ancestor::details[1]").evaluate("""element =>
                    new Promise(resolve => {
                        let previous = -1, stable = 0;
                        function check() {
                            const height = element.getBoundingClientRect().height;
                            stable = Math.abs(height - previous) < 0.1 ? stable + 1 : 0;
                            previous = height;
                            if (stable >= 8) resolve(true);
                            else requestAnimationFrame(check);
                        }
                        requestAnimationFrame(check);
                    })
                """)
                documentation.evaluate("""e => {
                    e.scrollIntoView({block:'start'});
                    e.closest('[data-testid="stMain"]').scrollBy(0, -80);
                }""")
                page.screenshot(path=str(tmp_path / f"documentation-{section}-{size}.png"))
                documentation.click()
                _settled(page)
        print(f"Synthetic long-label browser screenshots and dimensions: {tmp_path}")
    except Exception:
        page.screenshot(path=str(tmp_path / "layout-failure.png"), full_page=True)
        (tmp_path / "layout-failure.html").write_text(page.content(), encoding="utf-8")
        print(f"Synthetic layout failure artifacts: {tmp_path}")
        raise
    finally:
        context.close()


def test_native_candidate_selection_follows_main_only(
    browser: Browser, workspace_url: str, tmp_path: Path
) -> None:
    context = browser.new_context(viewport={"width": 1366, "height": 768}, accept_downloads=True)
    page = context.new_page()
    page.set_default_timeout(10000)
    try:
        _load(page, workspace_url, tmp_path, count=6)
        _expect_table_product(page, "000001")
        page.get_by_role("button", name="Næste kandidat", exact=True).click()
        _settled(page)
        _expect_product(page, "000002")
        _expect_table_product(page, "000002")
        # Product 1 is in the table, but its graph checkbox only adds an overlay.
        _toggle_alternative(page, "000001")
        _expect_product(page, "000002")
        _expect_table_product(page, "000002")
        _review_alternative(page, "000001")
        _expect_product(page, "000001")
        _expect_table_product(page, "000001")
        _show_review(page)
        page.get_by_role("textbox", name="Note", exact=True).fill("Synkroniseret hovedvare")
        page.get_by_role("button", name="Gem og næste kandidat", exact=True).click()
        _settled(page)
        _expect_product(page, "000002")
        _expect_table_product(page, "000002")
        page.get_by_role("button", name="Forrige kandidat", exact=True).click()
        _settled(page)
        _expect_product(page, "000001")
        _expect_table_product(page, "000001")
        _show_review(page, enabled=False)
        search = page.get_by_role("textbox", name="Søg i kandidatlisten", exact=True)
        search.fill("000003")
        search.press("Enter")
        _settled(page)
        _expect_product(page, "000001")
        _expect_table_product(page, None)
        _toggle_alternative(page, "000003")
        _expect_product(page, "000001")
        _expect_table_product(page, None)
        _open_from_table(page, "000003")
        _expect_table_product(page, "000003")
        search.fill("")
        search.press("Enter")
        _settled(page)
        _expect_table_product(page, "000003")
        page.get_by_text("Eksport og sessionskopi", exact=True).click()
        saved = _download(page, "Hent sessionskopi (JSON)", tmp_path)
        requested = page.get_by_role("spinbutton", name="Antal kandidater", exact=True)
        requested.fill("2")
        requested.press("Enter")
        _settled(page)
        _expect_product(page, "000003")
        _expect_table_product(page, None)
        grid = page.locator(".st-key-candidate_list").get_by_test_id("stDataFrame")
        expect(grid.locator('tbody [aria-colindex="2"]')).to_have_text(["000001", "000002"])
        grid.evaluate("""e => {
            e.scrollIntoView({block:'start'});
            e.closest('[data-testid="stMain"]').scrollBy(0, -80);
        }""")
        # Choose the nonfirst row while entering required-selection mode.
        grid.click(position={"x": 16, "y": 92})
        _settled(page)
        _expect_product(page, "000002")
        _expect_table_product(page, "000002")
        page.locator('input[type="file"]').nth(2).set_input_files(saved)
        page.get_by_role("button", name="Gendan vurderinger", exact=True).click()
        expect(page.get_by_text("Sessionskopien er gendannet", exact=False)).to_be_visible()
        _settled(page)
        _expect_product(page, "000003")
        _expect_table_product(page, "000003")
        page.locator(".st-key-candidate_list").screenshot(
            path=str(tmp_path / "native-main-000003-selected.png")
        )
    except Exception:
        page.screenshot(path=str(tmp_path / "selection-failure.png"), full_page=True)
        (tmp_path / "selection-failure.html").write_text(page.content(), encoding="utf-8")
        raise
    finally:
        context.close()


def test_decision_navigation(browser: Browser, workspace_url: str, tmp_path: Path) -> None:
    context = browser.new_context(viewport={"width": 1366, "height": 768}, accept_downloads=True)
    page = context.new_page()
    page.set_default_timeout(10000)
    try:
        _load(page, workspace_url, tmp_path, count=6)
        _show_review(page)
        # Human names are searched through the native candidate table.
        _open_from_table(page, "000002", query="Eksempelpakning 2")
        note = page.get_by_role("textbox", name="Note", exact=True)
        note.fill("Ugemt kladde til vare 2")
        note.press("Tab")
        expect(page.get_by_text("Ugemte ændringer.", exact=False)).to_be_visible()
        _settled(page)
        # Direct graph controls preserve the selected review owner and its unsaved draft.
        _toggle_alternative(page, "900001")
        for choice, periods in (("26 perioder", 26), ("52 perioder", 40), ("Alle", 40)):
            page.get_by_text(choice, exact=True).click()
            _settled(page)
            expect(page.get_by_role("radio", name=choice, exact=True)).to_be_checked()
            page.wait_for_function(
                "n => document.querySelector('.js-plotly-plot').data[0].x.length === n", arg=periods
            )
            _expect_product(page, "000002")
            expect(note).to_have_value("Ugemt kladde til vare 2")
            expect(_alternative(page, "900001")).to_be_checked()
        later = page.get_by_role("checkbox", name="Senere priser", exact=True)
        later.press("Space")
        _settled(page)
        expect(later).to_be_checked()
        _expect_product(page, "000002")
        expect(note).to_have_value("Ugemt kladde til vare 2")
        expect(_alternative(page, "900001")).to_be_checked()
        _select(page, "Vurdering", "Undersøg nærmere")
        page.get_by_text("Føj varen til min eksportliste", exact=True).click()
        _settled(page)
        _show_review(page, enabled=False)
        page.get_by_role("button", name="Næste kandidat", exact=True).click()
        _settled(page)
        _expect_product(page, "000003")
        page.get_by_text("26 perioder", exact=True).click()
        _settled(page)
        expect(page.locator(".st-key-review_editor")).to_have_count(0)
        page.get_by_role("button", name="Forrige kandidat", exact=True).click()
        _settled(page)
        _expect_product(page, "000002")
        expect(_alternative(page, "900001")).to_be_checked()
        page.get_by_text("Eksport og sessionskopi", exact=True).click()
        unsaved = json.loads(_download(page, "Hent sessionskopi (JSON)", tmp_path).read_bytes())
        assert "000002" not in unsaved["reviews"]
        page.get_by_text("Eksport og sessionskopi", exact=True).click()
        _show_review(page)
        expect(note).to_have_value("Ugemt kladde til vare 2")
        expect(page.get_by_role("combobox", name="Vurdering", exact=True)).to_have_value(
            "Undersøg nærmere"
        )
        expect(
            page.get_by_role("checkbox", name="Føj varen til min eksportliste", exact=True)
        ).to_be_checked()
        expect(page.get_by_text("Ugemte ændringer.", exact=False)).to_be_visible()
        page.get_by_text("Alle", exact=True).click()
        _settled(page)
        _toggle_alternative(page, "900001", checked=False)
        page.get_by_role("button", name="Næste kandidat", exact=True).click()
        _expect_product(page, "000003")
        _settled(page)
        expect(note).to_have_value("")
        page.get_by_role("button", name="Forrige kandidat", exact=True).click()
        _expect_product(page, "000002")
        _settled(page)
        expect(note).to_have_value("Ugemt kladde til vare 2")
        page.get_by_role("button", name="Gem og næste kandidat", exact=True).click()
        _expect_product(page, "000003")
        _settled(page)
        expect(page.get_by_role("tab")).to_have_count(0)
        expect(page.get_by_test_id("stPlotlyChart").first).to_be_visible()
        _open_from_table(page, "000002")
        expect(note).to_have_value("Ugemt kladde til vare 2")
        expect(page.get_by_text("Ugemte ændringer.", exact=False)).to_have_count(0)
        # Comparison only toggles traces; the candidate and its draft retain ownership.
        _toggle_alternative(page, "900001")
        _expect_product(page, "000002")
        expect(page.locator(".st-key-alternative_000002_900001")).to_contain_text(
            "Opb. H · Firma ikke oplyst"
        )
        expect(page.locator(".st-key-open_alternative_000002_900001")).to_have_count(0)
        page.wait_for_function(
            "() => {const p=document.querySelector('.js-plotly-plot');"
            "return p && p.data.length===2 && p.data[0].meta==='000002' "
            "&& p.data[1].meta==='900001' && p.data[0].line.dash!==p.data[1].line.dash "
            "&& p.data[0].marker.symbol!==p.data[1].marker.symbol}"
        )
        _toggle_alternative(page, "900001", checked=False)
        page.wait_for_function("() => document.querySelector('.js-plotly-plot').data.length===1")
        _toggle_alternative(page, "900001")
        expect(_alternative(page, "900002")).to_be_disabled()
        expect(page.locator(".st-key-alternative_000002_900002")).to_contain_text("Mangler")
        page.get_by_role("button", name="Help for 900002", exact=True).hover()
        expect(page.get_by_role("tooltip")).to_contain_text(
            "Historik mangler; ingen kurve kan vises."
        )
        page.mouse.move(0, 0)
        _expect_product(page, "000002")
        expect(page.locator(".st-key-open_alternative_000002_900002")).to_have_count(0)
        _review_alternative(page, "900004")
        _expect_product(page, "900004")
        expect(
            page.get_by_text("Historisk prisniveau kan ikke sammenlignes", exact=False)
        ).to_be_visible()
        expect(note).to_have_value("")
        _open_from_table(page, "000002")
        expect(note).to_have_value("Ugemt kladde til vare 2")
        expect(_alternative(page, "900001")).to_be_checked()
        _toggle_alternative(page, "000001")
        page.locator(".st-key-price_workspace").get_by_text(
            "Substitutionsalternativer", exact=False
        ).evaluate("""e => {
            e.scrollIntoView({block:'center'});
        }""")
        page.screenshot(path=str(tmp_path / "two-selected-alternatives-laptop.png"))
        _review_alternative(page, "000001")
        _expect_product(page, "000001")
        expect(note).to_have_value("")
        expect(_alternative(page, "900001")).not_to_be_checked()
        page.get_by_text("Eksport og sessionskopi", exact=True).click()
        saved = _download(page, "Hent sessionskopi (JSON)", tmp_path)
        payload = json.loads(saved.read_bytes())
        assert payload["analysis_id"] == unsaved["analysis_id"]
        assert payload["reviews"]["000002"]["note"] == "Ugemt kladde til vare 2"
        assert "000001" not in payload["reviews"]
        assert "900004" not in payload["reviews"]
    except Exception:
        page.screenshot(path=str(tmp_path / "decision-failure.png"), full_page=True)
        (tmp_path / "decision-failure.html").write_text(page.content(), encoding="utf-8")
        print(f"Synthetic decision artifacts: {tmp_path}")
        raise
    finally:
        context.close()


def test_group_checkbox_limit(browser: Browser, workspace_url: str, tmp_path: Path) -> None:
    context = browser.new_context(viewport={"width": 1366, "height": 768})
    page = context.new_page()
    page.set_default_timeout(10000)
    try:
        _load(page, workspace_url, tmp_path, count=9, large_group=True, descriptive_labels=True)
        rows = page.locator('.st-key-price_workspace [class*="st-key-alternative_"]')
        expect(rows).to_have_count(12)
        expect(page.get_by_text(re.compile(r"^Vis flere alternativer"))).to_have_count(0)
        expected_ids = [f"{i:06d}" for i in range(2, 10)] + ["900001", "900002", "900003", "900004"]
        for row, pid in zip(rows.all(), expected_ids, strict=True):
            expect(row.get_by_role("checkbox")).to_have_accessible_name(pid)
            text = row.get_by_test_id("stText")
            company = (
                "Firma ikke oplyst" if pid in ("900002", "900003", "900004") else f"Firma {pid}"
            )
            expect(text.nth(0)).to_have_text(f"Opb. H · {company}" if pid == "900001" else company)
            expect(text.nth(1)).to_have_text(
                "A → A" if pid == "900003" else "B → C" if pid == "900004" else "A → C"
            )
            expect(row.locator("[data-testid='stMarkdownContainer'] strong")).to_have_text(
                "Mangler" if pid in ("900002", "900003", "900004") else "100,00"
            )
        # Exactly one control per alternative, including disabled missing-history rows.
        expect(_alternative(page, "900002")).to_be_disabled()
        expect(page.get_by_role("checkbox", name=re.compile(r"^\d{6}$"))).to_have_count(12)
        expect(_alternative(page, "000009")).to_be_enabled()
        for pid in ("000002", "000003", "000004", "000005", "000006"):
            _toggle_alternative(page, pid)
        _settled(page)
        expect(_alternative(page, "000007")).to_be_disabled()
        expect(_alternative(page, "000002")).to_be_enabled()
        page.wait_for_function("() => document.querySelector('.js-plotly-plot').data.length===6")
        assert page.locator(".js-plotly-plot").evaluate(
            "plot => plot.data.map(trace => trace.name)"
        ) == [f"{i:06d} · Firma {i:06d}" for i in range(1, 7)]
        _unified_hover_evidence(page, tmp_path)
        page.keyboard.press("Escape")
        _linked_layout_evidence(page, tmp_path)
        _group_details_evidence(page, tmp_path, count=13)
        _toggle_alternative(page, "000002", checked=False)
        _settled(page)
        expect(_alternative(page, "000007")).to_be_enabled()
        _toggle_alternative(page, "000007")
        _settled(page)
        page.wait_for_function(
            "() => {const ids=document.querySelector('.js-plotly-plot').data.map(t=>t.meta);"
            "return ids.length===6 && ids.includes('000007') && !ids.includes('000002')}"
        )
        page.screenshot(path=str(tmp_path / "five-alternatives-laptop.png"), full_page=True)
        varied_context = browser.new_context(viewport={"width": 1366, "height": 768})
        try:
            varied = varied_context.new_page()
            varied.set_default_timeout(10000)
            varied_path = tmp_path / "varied-prices"
            varied_path.mkdir()
            _load(
                varied,
                workspace_url,
                varied_path,
                count=9,
                large_group=True,
                descriptive_labels=True,
                varied_prices=True,
            )
            _open_from_table(varied, "000001")
            for product_id in ("000002", "000003", "000004", "000005", "000006"):
                _toggle_alternative(varied, product_id)
            varied.wait_for_function(
                "() => document.querySelector('.js-plotly-plot').data.length===6"
            )
            sample = varied.locator(".js-plotly-plot").evaluate("""plot => plot.data.map(t => ({
                id:t.meta, x:t.x, y:t.y, line:t.line, marker:t.marker,
            }))""")
            assert len({tuple(trace["y"]) for trace in sample}) == 6
            (varied_path / "synthetic-line-styles.json").write_text(
                json.dumps(sample, indent=2), encoding="utf-8"
            )
            _linked_layout_evidence(varied, varied_path)
        finally:
            varied_context.close()
    finally:
        context.close()


def test_real_workflow(browser: Browser, workspace_url: str, tmp_path: Path) -> None:
    context = browser.new_context(viewport={"width": 1366, "height": 768}, accept_downloads=True)
    page = context.new_page()
    page.set_default_timeout(10000)
    try:
        _load(page, workspace_url, tmp_path)
        _show_review(page)
        page.get_by_role("button", name="Filtre", exact=True).click()
        expect(
            page.get_by_role("combobox", name="Vis kun denne substitutionsgruppe", exact=True)
        ).to_be_visible()
        expect(
            page.get_by_role("combobox", name="Vis kun denne vurdering", exact=True)
        ).to_be_visible()
        page.keyboard.press("Escape")
        requested = page.get_by_role("spinbutton", name="Antal kandidater", exact=True)
        requested.fill("80")
        requested.press("Enter")
        expect(page.get_by_text("Viser 80 af 80 ønskede kandidater", exact=False)).to_be_visible()
        # Glide's supported native dataframe is canvas-rendered. Its first header/row
        # coordinates are measured from the visible grid, never fixed screen positions.
        grid = page.locator('[data-testid="stDataFrame"]:visible').first
        grid.scroll_into_view_if_needed()
        canvas = grid.locator("canvas").first
        expect(canvas).to_be_visible()
        box = canvas.bounding_box()
        assert box is not None
        grid.click(position={"x": 16, "y": 56})
        _expect_product(page, "000001")
        _expect_table_product(page, "000001")
        _settled(page)
        page.get_by_role("button", name="Næste kandidat", exact=True).click()
        _expect_product(page, "000002")
        _expect_table_product(page, "000002")
        grid.click(position={"x": 16, "y": 56})
        _expect_product(page, "000001")
        _settled(page)
        # Native header sorting changes visible rows without changing canonical ranks.
        grid.evaluate("""e => {
            e.scrollIntoView({block:'start'});
            e.closest('[data-testid="stMain"]').scrollBy(0, -80);
        }""")
        grid.click(position={"x": 65, "y": 15})
        ascending = page.get_by_role("menuitem", name="Sort ascending", exact=True)
        if ascending.is_visible():
            ascending.click()
        expect(grid.locator("th").first).to_have_text("↑ Prioritet")
        grid.click(position={"x": 65, "y": 15})
        descending = page.get_by_role("menuitem", name="Sort descending", exact=True)
        if descending.is_visible():
            descending.click()
        expect(grid.locator("th").first).to_have_text("↓ Prioritet")
        grid.click(position={"x": 16, "y": 56})
        _expect_product(page, "000080")
        _expect_table_product(page, "000080")
        page.get_by_role("button", name="Forrige kandidat", exact=True).click()
        _settled(page)
        _expect_product(page, "000079")
        _expect_table_product(page, "000079")
        expect(grid.locator("th").first).to_have_text("↓ Prioritet")
        page.get_by_role("button", name="Næste kandidat", exact=True).click()
        _settled(page)
        _expect_product(page, "000080")
        _expect_table_product(page, "000080")
        expect(grid.locator("th").first).to_have_text("↓ Prioritet")
        expect(
            page.get_by_test_id("stText").filter(
                has_text=re.compile(r"^000080 · Eksempelpakning 80$")
            )
        ).to_be_visible()
        _settled(page)
        page.get_by_role("textbox", name="Note", exact=True).fill("=Syntetisk gennemgang")
        page.get_by_text("Føj varen til min eksportliste", exact=True).click()
        _select(page, "Vurdering", "Interessant")
        page.get_by_role("button", name="Gem vurdering", exact=True).click()
        expect(page.get_by_text("Vurderingen er gemt", exact=False)).to_be_visible()
        requested.fill("17")
        requested.press("Enter")
        expect(page.get_by_text("Viser 17 af 17 ønskede kandidater", exact=False)).to_be_visible()
        _expect_product(page, "000080")
        _expect_table_product(page, None)
        expect(page.get_by_role("textbox", name="Note", exact=True)).to_have_value(
            "=Syntetisk gennemgang"
        )
        expect(
            page.get_by_text("uden for den aktuelle prioriteringskø", exact=False)
        ).to_be_visible()
        search = page.get_by_role("textbox", name="Søg i kandidatlisten", exact=True)
        search.fill("000080")
        search.press("Enter")
        expect(page.get_by_text("Ingen produkter i denne visning", exact=False)).to_be_visible()
        search.fill("000002")
        search.press("Enter")
        expect(page.get_by_text("1 synlige efter søgning og filtre.", exact=False)).to_be_visible()
        _settled(page)
        expect(page.locator(".st-key-candidate_list").get_by_test_id("stDataFrame")).to_be_visible()
        _expect_product(page, "000080")
        expect(page.get_by_role("button", name="Forrige kandidat", exact=True)).to_be_disabled()
        expect(page.get_by_role("button", name="Næste kandidat", exact=True)).to_be_disabled()
        expect(page.get_by_role("textbox", name="Note", exact=True)).to_have_value(
            "=Syntetisk gennemgang"
        )
        search.fill("")
        search.press("Enter")
        _select(page, "Hvilke kandidater vil du se?", "Alle prioriterbare")
        expect(page.get_by_text("Alle 120 prioriterbare kandidater.", exact=False)).to_be_visible()
        _settled(page)
        page.get_by_text("26 perioder", exact=True).click()
        expect(page.get_by_role("radio", name="26 perioder", exact=True)).to_be_checked()
        page.get_by_role("checkbox", name="Senere priser", exact=True).press("Space")
        _open_from_table(page, "000001")
        expect(
            page.get_by_test_id("stText").filter(
                has_text=re.compile(r"^000001 · Eksempelpakning 1$")
            )
        ).to_be_visible()
        page.wait_for_function(
            "() => {const p=document.querySelector('.js-plotly-plot');"
            "return p && p.data[0].meta==='000001' && p.data[0].x.length===26}"
        )
        _settled(page)
        _toggle_alternative(page, "900001")
        expect(page.locator('[data-testid="stPlotlyChart"]').first).to_be_visible()
        page.wait_for_function(
            "() => {const p=document.querySelector('.js-plotly-plot');"
            "return p && p.data.length===2 && p.data[0].x.length===26 "
            "&& p.data[1].meta==='900001'}"
        )
        page.get_by_role("button", name="Gruppedetaljer", exact=True).click()
        expect(
            page.get_by_text("Gruppefuldstændighed er ikke verificeret", exact=False).last
        ).to_be_visible()
        expect(page.locator('[class*="st-key-group_detail_"]')).to_have_count(7)
        page.keyboard.press("Escape")
        _select(page, "Hvilke kandidater vil du se?", "Prioriteret liste")
        requested.fill("80")
        requested.press("Enter")
        page.get_by_text("Eksport og sessionskopi", exact=True).click()
        page.get_by_role("button", name="Klargør Excel", exact=True).click()
        workbook_path = _download(page, "Hent Excel", tmp_path)
        workbook = load_workbook(BytesIO(workbook_path.read_bytes()))
        assert workbook.sheetnames[0] == "Indkøbsoverblik"
        rows = list(workbook["Kandidater"].iter_rows(min_row=2, values_only=True))
        assert [row[1] for row in rows] == [f"{i:06d}" for i in range(1, 81)]
        assert rows[0][3] == 100 and isinstance(rows[0][3], (int, float))
        assert rows[79][25:27] == ("Interessant", "=Syntetisk gennemgang")
        assert workbook["Kandidater"]["AA81"].data_type == "s"
        saved = _download(page, "Hent sessionskopi (JSON)", tmp_path)
        payload = json.loads(saved.read_bytes())
        assert payload["queue"] == {"mode": "top_n", "top_n": 80}
        assert payload["reviews"]["000080"]["note"] == "=Syntetisk gennemgang"
        note = page.get_by_role("textbox", name="Note", exact=True)
        note.fill("Midlertidig ugemt kladde før gendannelse")
        note.press("Tab")
        expect(page.get_by_text("Ugemte ændringer.", exact=False)).to_be_visible()
        requested.fill("17")
        requested.press("Enter")
        page.locator('input[type="file"]').nth(2).set_input_files(saved)
        page.get_by_role("button", name="Gendan vurderinger", exact=True).click()
        expect(page.get_by_text("Sessionskopien er gendannet", exact=False)).to_be_visible()
        expect(requested).to_have_value("80")
        expect(note).to_have_value("")
        expect(page.get_by_text("Ugemte ændringer.", exact=False)).to_have_count(0)
        page.wait_for_function(
            "() => {const p=document.querySelector('.js-plotly-plot');"
            "return p && p.data.length===2 && p.data[1].meta==='900001'}"
        )
        _open_from_table(page, "000080")
        expect(note).to_have_value("=Syntetisk gennemgang")
        _open_from_table(page, "000001")
        expect(note).to_have_value("")
        page.wait_for_function(
            "() => {const p=document.querySelector('.js-plotly-plot');"
            "return p && p.data.length===2 && p.data[0].meta==='000001' "
            "&& p.data[1].meta==='900001'}"
        )
        _linked_layout_evidence(page, tmp_path)
        for size, width, height in (("laptop", 1366, 768), ("desktop", 1920, 1080)):
            page.set_viewport_size({"width": width, "height": height})
            for section, heading in (
                ("workspace", "Redbird · indkøbsgennemgang"),
                ("table", "1. Vælg en kandidat"),
                ("detail", "2. Sammenlign priser"),
            ):
                page.get_by_role("heading", name=heading, exact=True).evaluate(
                    "e => e.scrollIntoView({block:'start'})"
                )
                page.screenshot(path=str(tmp_path / f"{section}-{size}.png"))
            page.get_by_test_id("stPlotlyChart").first.evaluate(
                "e => e.scrollIntoView({block:'start'})"
            )
            page.screenshot(path=str(tmp_path / f"chart-{size}.png"))
        page.set_viewport_size({"width": 1366, "height": 768})
        # A failed replacement may not overwrite the loaded sources or saved notes.
        page.get_by_text("Datagrundlag og indstillinger", exact=True).click()
        page.get_by_text("Jeg accepterer at erstatte den aktive vurdering", exact=False).click()
        page.locator('input[type="file"]').nth(0).set_input_files(
            {
                "name": "synthetic-broken.xlsx",
                "mimeType": "application/octet-stream",
                "buffer": b"invalid workbook",
            }
        )
        # Upload completion and the server rerun are separate browser events. The
        # pending-input notice proves Streamlit has consumed the changed source.
        expect(
            page.get_by_role("button", name="Remove synthetic-broken.xlsx", exact=True)
        ).to_be_visible()
        expect(page.get_by_text("Der er ændrede input.", exact=False)).to_be_visible()
        _settled(page)
        page.get_by_role("button", name="Anvend indstillinger", exact=True).click()
        expect(page.get_by_text("tidligere, sammenhængende analyse", exact=False)).to_be_visible()
        _open_from_table(page, "000080")
        expect(page.get_by_role("textbox", name="Note", exact=True)).to_have_value(
            "=Syntetisk gennemgang"
        )
        page.get_by_role("heading", name="Redbird · indkøbsgennemgang", exact=True).evaluate(
            "e => e.scrollIntoView({block:'start'})"
        )
        page.screenshot(path=str(tmp_path / "error-laptop.png"))
        print(f"Synthetic browser screenshots: {tmp_path}")
        other = browser.new_context(viewport={"width": 1366, "height": 768})
        try:
            isolated = other.new_page()
            isolated.goto(workspace_url)
            expect(
                isolated.get_by_text("Upload Cito-rapport og LMS-prishistorik", exact=False)
            ).to_be_visible(timeout=30000)
            expect(isolated.get_by_role("textbox", name="Note", exact=True)).to_have_count(0)
            expect(isolated.get_by_text("synthetic-cito.xlsx", exact=False)).to_have_count(0)
            _load(isolated, workspace_url, tmp_path, count=3)
            _show_review(isolated)
            expect(isolated.get_by_role("textbox", name="Note", exact=True)).to_have_value("")
            isolated.get_by_role("textbox", name="Note", exact=True).fill("Anden session")
            isolated.get_by_role("button", name="Gem vurdering", exact=True).click()
            expect(isolated.get_by_text("Vurderingen er gemt", exact=False)).to_be_visible()
            _open_from_table(page, "000001")
            expect(page.get_by_role("textbox", name="Note", exact=True)).to_have_value("")
            _open_from_table(page, "000080")
            expect(page.get_by_role("textbox", name="Note", exact=True)).to_have_value(
                "=Syntetisk gennemgang"
            )
            expect(page.get_by_text("120 kan prioriteres", exact=False).first).to_be_visible()
        finally:
            other.close()
    except Exception:
        page.screenshot(path=str(tmp_path / "failure.png"), full_page=True)
        (tmp_path / "failure.html").write_text(page.content(), encoding="utf-8")
        print(f"Synthetic browser failure artifacts: {tmp_path}")
        raise
    finally:
        context.close()
